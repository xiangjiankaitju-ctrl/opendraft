#!/usr/bin/env python3
"""Regression tests for production DOCX export formatting pipeline."""

import os
import json
import re
import shutil
import sys
import zipfile
from pathlib import Path

import pytest


PROJECT_ROOT = Path(__file__).resolve().parent.parent
ENGINE_ROOT = PROJECT_ROOT / "engine"
sys.path.insert(0, str(ENGINE_ROOT))

from utils.document_ast import build_document_structure_manifest, normalize_document_markdown
from utils.docx_export_pipeline import clean_language_residuals, normalize_docx_language, preprocess_markdown_for_docx, select_reference_template
from utils.export_professional import _ensure_yaml_title_schema, export_docx
from utils.final_artifact_contract import (
    clean_citation_residuals,
    final_artifact_validation,
    normalize_pagebreaks,
)
from utils.outline_contract import extract_docx_outline, extract_markdown_outline, outline_signature


ZH_SAMPLE = """---
title: 人工智能背景下就业观念变迁及其社会影响研究
language: zh
date: 2026-05-02
---

人工智能背景下就业观念变迁及其社会影响研究

# Table of Contents

## 1. 摘要

**Research Problem and Approach:** 本文梳理人工智能背景下就业观念变迁。

**Methodology and Findings:** 文献显示就业评价标准发生变化。

**Key Contributions:** 提供社会影响分析框架。

**Implications:** 为教育与劳动政策提供参考。

**Keywords:** 人工智能；就业观念；社会影响

\\newpage

# · 1. Introduction

正文段落。

## 1.1 Theoretical Framework

表 9 错误旧编号

| 维度 | 影响路径 |
| --- | --- |
| 技能 | 岗位重组 |

# 5. References

Https://doi.org/10.1000/example
"""


def test_front_matter_title_key_preserved():
    repaired = _ensure_yaml_title_schema(
        '---\n: "论文题名"\nauthor: "OpenDraft AI"\nlanguage: "zh"\ndate: "May 2026"\n---\n\n# 摘要\n正文',
        "fallback_title",
    )

    assert '\ntitle: "论文题名"\n' in repaired
    assert '\n: "论文题名"' not in repaired


def test_chinese_language_cleanup_uses_natural_section_references():
    text = (
        "文献综述（section 2.1）指出，离线编程技术具有价值。\n"
        "The findings presented in section 2.3 show additional constraints.\n"
        "文献综述（相关章节）揭示了路径规划不足。\n"
    )

    cleaned, fixed = clean_language_residuals(text, "zh")

    assert "section 2.1" not in cleaned
    assert "section 2.3" not in cleaned
    assert "相关章节" not in cleaned
    assert "文献综述（" not in cleaned
    assert "前文综述指出" in cleaned
    assert "第2.3节" in cleaned
    assert "前文综述揭示" in cleaned
    assert "section_or_phrase_residual" in fixed


def test_yaml_title_schema_repairs_keyless_title():
    test_front_matter_title_key_preserved()


def test_pagebreak_normalized():
    variants = "\n".join(
        [
            "<!—— PAGEBREAK ——>",
            "<!— PAGEBREAK —>",
            "<!-- PAGEBREAK -->",
            "<!--PAGEBREAK-->",
            r"\newpage",
            "/newpage",
            "ewpage",
            "newpage",
            "<!-- PAGEBREAK --><!-- PAGEBREAK -->",
        ]
    )

    normalized = normalize_pagebreaks(variants)
    assert "<!—— PAGEBREAK ——>" not in normalized
    assert "<!— PAGEBREAK —>" not in normalized
    assert r"\newpage" not in normalized
    assert "/newpage" not in normalized
    assert "ewpage" not in normalized
    assert "newpage" not in normalized.replace("PAGEBREAK", "")
    assert normalized.count("<!-- PAGEBREAK -->") == 1

    processed, _stats = preprocess_markdown_for_docx(
        f"---\ntitle: t\nlanguage: en\n---\n\n# 1. Introduction\n\n{variants}\n\nBody.",
        "en",
    )
    assert "PAGEBREAK" not in processed
    assert "```{=openxml}" in processed


def test_docx_preprocess_protects_markdown_technical_tokens_and_dashes():
    md = """---
title: 智能清洁机器人路径规划研究
language: zh
date: 2026-05-08
---

# 摘要

本文比较D*算法、CCD*覆盖、D* Lite、A*、C++、C#与<100 kHz阈值。

----包括连续破折号。

# 1. 引言

正文。
"""
    processed, stats = preprocess_markdown_for_docx(md, "zh")

    assert "D\\*算法" in processed
    assert "CCD\\*" in processed
    assert "D\\* Lite" in processed
    assert "A\\*" in processed
    assert "C++" in processed
    assert "C#" in processed
    assert "<100 kHz" in processed
    assert "——包括" in processed
    assert "—-" not in processed
    assert stats.technical_tokens_protected is True
    assert stats.damaged_tokens == []


def test_docx_cover_does_not_fallback_to_abstract_heading(monkeypatch, tmp_path):
    from docx import Document
    from utils import docx_post_processor as post

    docx_path = tmp_path / "cover.docx"
    doc = Document()
    doc.add_heading("摘要", level=1)
    doc.add_paragraph("摘要正文")
    doc.add_heading("1. 引言", level=1)
    doc.add_heading("1.1 背景", level=2)
    doc.add_heading("2. 文献综述", level=1)
    doc.save(docx_path)

    monkeypatch.setattr(post, "_update_fields_with_libreoffice", lambda _path, stats: False)
    stats = post.insert_academic_structure(
        docx_path,
        options={"language": "zh", "title": "论文题名", "date": "2026-05-08", "filename_stem": "fallback"},
    )
    texts = [p.text.strip() for p in Document(docx_path).paragraphs if p.text.strip()]

    assert texts[0] == "论文题名"
    assert texts[1] == "2026年5月"
    assert "摘要" in texts
    assert stats["cover_title"] == "论文题名"
    assert stats["title_from_abstract_heading"] is False
    assert stats["toc_field_inserted"] is True
    assert stats["toc_refreshed"] is False
    assert stats["manual_update_required"] is True
    assert stats["format_status"] == "needs_manual_toc_update"


def test_cover_title_not_abstract(monkeypatch, tmp_path):
    test_docx_cover_does_not_fallback_to_abstract_heading(monkeypatch, tmp_path)


def test_citation_residual_cleanup():
    dirty = "中文{{cite_001}} 继续{(Raj & Kos, 2022)} 以及{{Chen, 2025}}。"
    cleaned = clean_citation_residuals(dirty, "zh")

    assert "cite_001" not in cleaned
    assert "{(" not in cleaned
    assert "{{" not in cleaned
    assert "（Raj & Kos, 2022）" in cleaned
    assert "（Chen, 2025）" in cleaned

    processed, stats = preprocess_markdown_for_docx(
        f"---\ntitle: t\nlanguage: zh\n---\n\n# 摘要\n\n{dirty}\n\n# 1. 引言\n\n正文。",
        "zh",
    )
    assert "cite_001" not in processed
    assert "{(" not in processed
    assert "{{" not in processed
    assert stats.citation_residue_repaired is True


ZH_DRONE_SAMPLE = """---
title: 基于无人机遥感与深度学习的农作物病虫害识别研究
language: zh
date: 2026-05-02
---

基于无人机遥感与深度学习的农作物病虫害识别研究

Title

Document Type

Generated by

Disclaimer

Table of Contents

Right-click and update field to refresh the table of contents.

# 1. 摘要

本文研究无人机遥感与深度学习在农作物病虫害识别中的应用。

**关键词：** 无人机遥感；深度学习；病虫害识别

\\newpage

# · 1. 引言

无人机遥感为农田病虫害监测提供了高时空分辨率数据。

# · 2. 正文

表1：不同病害识别技术对比

| 技术类型 | 代表算法 | 应用优势 | 局限性 |
| --- | --- | --- | --- |
| 传统机器学习 | SVM | 可解释性较强 | 特征依赖人工设计 |
| 深度学习 | CNN | 自动提取特征 | 需要较多标注数据 |

Table 2. Table 2 识别性能比较

| Performance | Accuracy |
| --- | --- |
| CNN | 92% |

## 2.1 方法论框架

方法包括影像采集、标注、模型训练与验证。

## 2.2 分析与结果

结果表明深度学习模型具有较高识别精度。

## 2.3 讨论

模型泛化能力仍需增强。

# · 3. 结论

该方法可提升病虫害识别效率{(Wang et al., 2023)}{(萧 男 et al., 2025)}。

# 4. References

Wang et al. (2023). Example.
"""


EN_SAMPLE = """---
title: Microstructural Evolution and Performance Control of Metallic Materials in Additive Manufacturing
language: en
date: 2026-05-02
---

## Abstract

**Research Problem and Approach:** This draft reviews metallic materials in additive manufacturing.

**Keywords:** Additive Manufacturing; Laser Powder Bed Fusion

newpage

# 1. Introduction

Cooling rates exceeding $10^6$ K/s influence thermal gradients ($G$), growth velocity ($R$), and the resulting $G/R$ ratio.

Table 8. Old duplicate number

| Microstructural Outcome | Applicable Material Classes | Representative Study | Documented impacts |
| --- | --- | --- | --- |
| Columnar grains | Nickel alloys and steels | Example et al. | Strength anisotropy |

## 1.1 Background

Text.

# References

https://doi.org/10.1000/example
"""


def test_preprocess_chinese_localizes_and_cleans_docx_markdown():
    cleaned, stats = preprocess_markdown_for_docx(ZH_SAMPLE, "zh")

    assert stats.language == "zh"
    assert stats.tables_processed == 1
    assert stats.captions_generated == 1
    assert "Table of Contents" not in cleaned
    assert "人工智能背景下就业观念变迁及其社会影响研究" not in cleaned
    assert "Research Problem and Approach" not in cleaned
    assert "研究问题与方法" in cleaned
    assert "关键词：" in cleaned
    assert "```{=openxml}" in cleaned
    assert "ewpage" not in cleaned
    assert "\\newpage" not in cleaned
    assert "Https://doi.org" not in cleaned
    assert "https://doi.org/10.1000/example" in cleaned
    assert "表1：错误旧编号" in cleaned
    assert "表1\n\n|" not in cleaned
    assert "# 摘要" in cleaned
    assert "# 1. 引言" in cleaned
    assert "# · 1. 引言" not in cleaned
    assert "# 参考文献" in cleaned


def test_preprocess_english_preserves_labels_and_math_readability():
    cleaned, stats = preprocess_markdown_for_docx(EN_SAMPLE, "en")

    assert stats.language == "en"
    assert "Table 1." in cleaned
    assert "10⁶ K/s" in cleaned
    assert "G/R ratio" in cleaned
    assert "Research Problem and Approach" in cleaned
    assert "```{=openxml}" in cleaned
    assert "newpage" not in cleaned


def test_preprocess_chinese_removes_duplicate_heading_ordinals():
    sample = """---
title: 双重编号测试
language: zh
---

# 1. 引言

## 1.1 一、研究背景

正文。

## 1.2 二、问题提出

正文。
"""

    cleaned, _stats = preprocess_markdown_for_docx(sample, "zh")

    assert "## 1.1 研究背景" in cleaned
    assert "## 1.2 问题提出" in cleaned
    assert "一、研究背景" not in cleaned
    assert "二、问题提出" not in cleaned


def test_docx_preprocess_does_not_change_outline():
    sample = """---
title: 结构冻结测试
language: zh
---

# 1. 引言
正文。

# 2. 文献综述
引导段。

## 2.1 理论基础
正文。

### 2.1.1 机制分析
正文。

# 3. 研究方法
正文。

# 4. 分析结果
正文。

# 5. 讨论
正文。

# 6. 结论
正文。
"""

    cleaned, _stats = preprocess_markdown_for_docx(sample, "zh")

    assert outline_signature(extract_markdown_outline(sample)) == outline_signature(extract_markdown_outline(cleaned))
    assert cleaned.count("# 2. 文献综述") == 1
    assert cleaned.index("# 3. 研究方法") > cleaned.index("# 2. 文献综述")
    assert cleaned.index("# 3. 研究方法") < cleaned.index("# 5. 讨论")


def test_preprocess_detects_malformed_caption_pipe_table():
    malformed = """---
title: Bad table
language: zh
---

Table 1. Table 1 | 技术类型 | 代表算法 |
| 深度学习 | CNN |
"""
    with pytest.raises(ValueError, match="Malformed markdown table"):
        preprocess_markdown_for_docx(malformed, "zh")


def test_chinese_detection_tolerates_many_english_references():
    text = (
        "本文讨论无人机遥感与深度学习在农作物病虫害识别中的应用，并分析模型泛化能力与田间部署价值。" * 20
        + "\n".join(f"https://doi.org/10.1000/example{i} Wang et al. reference text" for i in range(30))
    )
    assert normalize_docx_language(None, text) == "zh"


@pytest.mark.parametrize("alias", ["zh", "zh-cn", "chinese", "cn", "中文"])
def test_docx_language_aliases_normalize_to_zh(alias):
    assert normalize_docx_language(alias, "") == "zh"


@pytest.mark.parametrize("alias", ["en", "english", "英文"])
def test_docx_language_aliases_normalize_to_en(alias):
    assert normalize_docx_language(alias, "") == "en"


def test_preprocess_chinese_drone_sample_removes_regression_residue():
    cleaned, stats = preprocess_markdown_for_docx(ZH_DRONE_SAMPLE, "zh")

    assert stats.tables_processed == 2
    assert stats.captions_generated == 2

    forbidden = [
        "基于无人机遥感与深度学习的农作物病虫害识别研究",
        "Table of Contents",
        "Right-click and update field",
        "\nTitle\n",
        "Document Type",
        "Generated by",
        "Disclaimer",
        "Table 1. Table 1",
        "Table 2. Table 2",
        "# 1. 摘要",
        "# · 1. 引言",
        "{(",
        ")}",
        "\\newpage",
        "ewpage",
    ]
    for marker in forbidden:
        assert marker not in cleaned

    required = [
        "# 摘要",
        "关键词",
        "# 1. 引言",
        "# 2. 文献综述",
        "# 3. 结论",
        "# 参考文献",
        "表1：不同病害识别技术对比",
        "表2：识别性能比较",
        "| 技术类型 | 代表算法 | 应用优势 | 局限性 |",
    ]
    for marker in required:
        assert marker in cleaned


def test_preprocess_chinese_repairs_braced_author_year_citations():
    sample = """---
title: 引用清理测试
language: zh
---

# 1. 引言

研究压力 {{Froholdt, 2018}} 仍在上升，韧性 {(Balcombe et al., 2019)} 也受关注。
"""

    cleaned, stats = preprocess_markdown_for_docx(sample, "zh")

    assert "{{Froholdt, 2018}}" not in cleaned
    assert "{(Balcombe et al., 2019)}" not in cleaned
    assert "压力（Froholdt, 2018）" in cleaned
    assert "韧性（Balcombe et al., 2019）" in cleaned
    assert stats.citation_residue_repaired is True


def test_preprocess_converts_table_caption_headings_to_plain_captions():
    sample = """---
title: 表题标题测试
language: zh
---

# 1. 正文

### 1.3.1. 表1 技术路线比较
| 维度 | 内容 |
| --- | --- |
| 路线 | 对比 |
"""

    cleaned, stats = preprocess_markdown_for_docx(sample, "zh")

    assert stats.markdown_tables_detected == 1
    assert stats.caption_headings_normalized == 1
    assert stats.captions_generated == 1
    assert "### 1.3.1. 表1" not in cleaned
    assert "表1：技术路线比较\n\n| 维度 | 内容 |" in cleaned


def test_preprocess_english_table_caption_heading_to_plain_caption():
    sample = """---
title: Caption heading test
language: en
---

# 1. Body

### 1.2.1. Table 3 Evaluation results
| A | B |
| --- | --- |
| x | y |
"""

    cleaned, stats = preprocess_markdown_for_docx(sample, "en")

    assert stats.markdown_tables_detected == 1
    assert stats.caption_headings_normalized == 1
    assert "### 1.2.1. Table 3" not in cleaned
    assert "Table 1. Evaluation results\n\n| A | B |" in cleaned


def test_preprocess_does_not_turn_table_reference_sentence_into_caption():
    sample = """---
title: 表格引用测试
language: zh
---

# 1. 正文

表 2 总结了主流技术方案的核心参数与应用场景对比。

| 技术路径 | 核心算法 |
| --- | --- |
| 图像增强 + 检测 | Extended ESRGAN + SSD |

*表 5：主流技术方案对比。*
"""

    cleaned, stats = preprocess_markdown_for_docx(sample, "zh")

    assert stats.markdown_tables_detected == 1
    assert stats.captions_generated == 1
    assert "表1 总结了主流技术方案的核心参数与应用场景对比。" in cleaned
    assert "表1：主流技术方案对比。" in cleaned
    assert "表1\n\n| 技术路径" not in cleaned


def test_preprocess_global_table_numbering_handles_section_numbered_caption_and_refs():
    sample = """---
title: 表格编号测试
language: zh
---

# 摘要
摘要正文。

<!-- PAGEBREAK --><!-- PAGEBREAK -->

# 1. 引言

正文引用表 6.1 总结了主要研究发现。

表 6.1：主要研究发现与管理启示摘要

| 发现 | 启示 |
| --- | --- |
| A | B |
"""

    cleaned, stats = preprocess_markdown_for_docx(sample, "zh")

    assert stats.markdown_tables_detected == 1
    assert "表1：主要研究发现与管理启示摘要" in cleaned
    assert "表1：1：" not in cleaned
    assert "正文引用表1 总结了主要研究发现。" in cleaned
    assert cleaned.count("```{=openxml}") == 1
    assert "<!-- PAGEBREAK -->" not in cleaned


@pytest.mark.parametrize(
    "bad_text",
    [
        "thermal gradients (), growth velocity R",
        "cooling rates exceeding  K/s",
        "the resulting  ratio",
    ],
)
def test_preprocess_blocks_empty_math_variables(bad_text):
    with pytest.raises(ValueError):
        preprocess_markdown_for_docx(f"---\ntitle: t\n---\n\n{bad_text}", "en")


def test_preprocess_allows_low_risk_empty_parentheses():
    cleaned, _stats = preprocess_markdown_for_docx("This section uses empty parentheses () as an editorial marker.", "en")
    assert "()" in cleaned


def test_reference_templates_exist():
    assert select_reference_template("zh").exists()
    assert select_reference_template("en").exists()


def _style_xml(path: Path, style_id: str) -> str:
    with zipfile.ZipFile(path) as zf:
        xml = zf.read("word/styles.xml").decode("utf-8")
    match = re.search(
        rf'<w:style\b(?=[^>]*w:styleId="{re.escape(style_id)}")[\s\S]*?</w:style>',
        xml,
    )
    assert match, f"Missing style {style_id} in {path}"
    return match.group(0)


def test_chinese_reference_template_is_readable_style_source():
    template = select_reference_template("zh")
    with zipfile.ZipFile(template) as zf:
        styles_xml = zf.read("word/styles.xml").decode("utf-8")

    normal = _style_xml(template, "Normal")
    heading1 = _style_xml(template, "Heading1")
    heading2 = _style_xml(template, "Heading2")
    heading3 = _style_xml(template, "Heading3")
    caption = _style_xml(template, "Caption")

    assert "SimSun" not in styles_xml
    assert 'w:eastAsia="Noto Serif CJK SC"' in normal
    assert 'w:ascii="Times New Roman"' in normal
    assert 'w:sz w:val="24"' in normal
    assert 'w:line="360"' in normal
    assert 'w:after="60"' in normal
    assert 'w:firstLine="420"' in normal
    assert 'w:jc w:val="left"' in normal

    assert 'w:sz w:val="32"' in heading1
    assert 'w:before="240"' in heading1
    assert 'w:after="120"' in heading1
    assert 'w:firstLine="0"' in heading1
    assert 'w:sz w:val="28"' in heading2
    assert 'w:before="200"' in heading2
    assert 'w:after="80"' in heading2
    assert 'w:sz w:val="24"' in heading3
    assert 'w:before="120"' in heading3
    assert 'w:after="60"' in heading3
    assert 'w:jc w:val="left"' in caption


def test_english_reference_template_uses_left_aligned_body_and_headings():
    template = select_reference_template("en")
    normal = _style_xml(template, "Normal")
    heading1 = _style_xml(template, "Heading1")

    assert 'w:sz w:val="22"' in normal
    assert 'w:jc w:val="left"' in normal
    assert 'w:firstLine=' not in normal or 'w:firstLine="0"' in normal
    assert 'w:jc w:val="left"' in heading1


def test_export_docx_requires_pandoc_and_does_not_call_basic(monkeypatch, tmp_path):
    input_md = tmp_path / "draft.md"
    input_md.write_text(EN_SAMPLE, encoding="utf-8")

    import utils.export_professional as export_module

    original_which = shutil.which
    monkeypatch.setattr(shutil, "which", lambda cmd: None if cmd == "pandoc" else original_which(cmd))
    monkeypatch.setattr(
        export_module,
        "export_docx_basic",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(AssertionError("basic fallback called")),
    )

    assert export_docx(input_md, tmp_path / "out.docx") is False


def test_export_docx_requires_reference_template(monkeypatch, tmp_path):
    input_md = tmp_path / "draft.md"
    input_md.write_text(EN_SAMPLE, encoding="utf-8")

    import utils.export_professional as export_module

    monkeypatch.setattr(shutil, "which", lambda cmd: "/usr/bin/pandoc" if cmd == "pandoc" else None)
    monkeypatch.setattr(export_module, "select_reference_template", lambda _lang: tmp_path / "missing.docx")

    assert export_docx(input_md, tmp_path / "out.docx") is False


def _docx_xml(path: Path, name: str = "word/document.xml") -> str:
    with zipfile.ZipFile(path) as zf:
        return zf.read(name).decode("utf-8")


def _all_docx_xml(path: Path) -> str:
    with zipfile.ZipFile(path) as zf:
        return "\n".join(
            zf.read(name).decode("utf-8", errors="ignore")
            for name in zf.namelist()
            if name.startswith("word/") and name.endswith(".xml")
        )


def _ensure_paragraph_style(doc, name: str):
    from docx.enum.style import WD_STYLE_TYPE

    try:
        return doc.styles[name]
    except KeyError:
        return doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)


def test_docx_post_processor_creates_word_structures(tmp_path):
    docx = pytest.importorskip("docx")
    from utils.docx_post_processor import insert_academic_structure

    output = tmp_path / "sample.docx"
    doc = docx.Document()
    doc.add_heading("摘要", level=1)
    doc.add_paragraph("关键词：人工智能；就业观念")
    doc.add_heading("1. 引言", level=1)
    doc.add_heading("1.1 研究背景", level=2)
    doc.add_paragraph("正文段落。")
    doc.add_paragraph("表 9  旧表题")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "维度"
    table.cell(0, 1).text = "影响"
    table.cell(1, 0).text = "技能"
    table.cell(1, 1).text = "岗位"
    doc.add_heading("参考文献", level=1)
    doc.add_paragraph("作者. (2025). 题名. https://doi.org/10.1000/example")
    doc.save(output)

    stats = insert_academic_structure(
        output,
        options={"language": "zh", "title": "人工智能背景下就业观念变迁及其社会影响研究"},
    )

    assert stats["tables_processed"] == 1
    assert stats["captions_generated"] == 1
    assert stats["docx_tables_detected"] == 1

    xml = _docx_xml(output)
    assert "Right-click and update field" not in xml
    assert 'w:type="page"' in xml
    assert "tblHeader" in xml
    assert "cantSplit" in xml
    assert "tblW" in xml
    assert "tblGrid" in xml
    assert xml.count("gridCol") >= 2

    processed = docx.Document(output)
    texts = "\n".join(p.text for p in processed.paragraphs)
    assert "目录" in texts or 'TOC \\o "1-3"' in xml or "TOC \\o &quot;1-3&quot;" in xml
    assert "摘要" in texts
    assert "关键词" in texts
    assert "参考文献" in texts
    assert "Table of Contents" not in texts
    assert "Document Type" not in texts
    assert "Generated by" not in texts
    assert "Disclaimer" not in texts
    assert "ewpage" not in texts


@pytest.mark.skipif(shutil.which("pandoc") is None, reason="Pandoc is not installed")
def test_export_docx_generates_real_docx_when_pandoc_available(tmp_path):
    docx = pytest.importorskip("docx")
    input_md = tmp_path / "draft.md"
    output = tmp_path / "draft.docx"
    input_md.write_text(EN_SAMPLE, encoding="utf-8")

    assert export_docx(input_md, output)

    document = docx.Document(output)
    texts = "\n".join(p.text for p in document.paragraphs)
    assert "Table of Contents" in texts
    assert "Abstract" in texts
    assert "Keywords" in texts
    assert "References" in texts
    assert "newpage" not in texts
    assert len(document.tables) == 1


def test_docx_post_processor_does_not_create_empty_toc_or_page_numbers(tmp_path):
    docx = pytest.importorskip("docx")
    from utils.docx_post_processor import insert_academic_structure

    output = tmp_path / "no_toc.docx"
    doc = docx.Document()
    doc.add_heading("摘要", level=1)
    doc.add_paragraph("关键词：无人机遥感")
    doc.add_heading("1. 引言", level=1)
    doc.add_paragraph("正文。")
    doc.save(output)

    assert insert_academic_structure(output, options={"language": "zh", "title": "测试文档"})
    processed = docx.Document(output)
    texts = "\n".join(p.text for p in processed.paragraphs)
    assert "目录" not in texts
    assert "Table of Contents" not in texts
    assert "PAGE" not in _all_docx_xml(output)


def test_docx_post_processor_preserves_toc_field_when_refresh_fails(monkeypatch, tmp_path):
    docx = pytest.importorskip("docx")
    import utils.docx_post_processor as post

    output = tmp_path / "static_toc.docx"
    doc = docx.Document()
    doc.add_heading("1. Introduction", level=1)
    doc.add_heading("1.1 Background", level=2)
    doc.add_paragraph("Body.")
    doc.add_heading("2. Literature Review", level=1)
    doc.add_heading("2.1 Prior Work", level=2)
    doc.add_paragraph("Body.")
    doc.save(output)

    monkeypatch.setattr(post, "_update_fields_with_libreoffice", lambda _path, stats: False)

    stats = post.insert_academic_structure(output, options={"language": "en", "title": "Static TOC"})
    processed = docx.Document(output)
    texts = [p.text.strip() for p in processed.paragraphs if p.text.strip()]

    assert "Table of Contents" in texts
    assert "1. Introduction" in texts
    assert "1.1 Background" in texts
    assert "2. Literature Review" in texts
    assert any("TOC field inserted but automatic refresh failed" in warning for warning in stats["warnings"])
    assert 'TOC \\o "1-3"' in _all_docx_xml(output) or "TOC \\o &quot;1-3&quot;" in _all_docx_xml(output)
    assert texts.count("1. Introduction") == 1
    assert texts.count("1.1 Background") == 1
    assert stats["inspection"]["static_toc_detected"] is False


def test_no_static_toc_fallback(monkeypatch, tmp_path):
    test_docx_post_processor_preserves_toc_field_when_refresh_fails(monkeypatch, tmp_path)


def test_toc_field_inserted_depth_3(monkeypatch, tmp_path):
    docx = pytest.importorskip("docx")
    import utils.docx_post_processor as post

    output = tmp_path / "toc_depth.docx"
    doc = docx.Document()
    doc.add_heading("Abstract", level=1)
    doc.add_paragraph("Body.")
    doc.add_heading("1. Introduction", level=1)
    doc.add_heading("1.1 Background", level=2)
    doc.add_heading("1.1.1 Details", level=3)
    doc.add_heading("2. Literature Review", level=1)
    doc.add_heading("2.1 Prior Work", level=2)
    doc.save(output)
    monkeypatch.setattr(post, "_update_fields_with_libreoffice", lambda _path, stats: False)

    stats = post.insert_academic_structure(output, options={"language": "en", "title": "TOC Depth"})
    xml = _all_docx_xml(output)

    assert stats["toc_field_inserted"] is True
    assert 'TOC \\o "1-3"' in xml or "TOC \\o &quot;1-3&quot;" in xml


def test_toc_depth_3(monkeypatch, tmp_path):
    test_toc_field_inserted_depth_3(monkeypatch, tmp_path)


def test_toc_includes_heading3(monkeypatch, tmp_path):
    docx = pytest.importorskip("docx")
    import utils.docx_post_processor as post

    md = tmp_path / "heading3.md"
    md.write_text(
        """---
title: "Heading 3 TOC"
language: "zh"
date: "May 2026"
---

# 摘要
摘要正文。

# 1. 引言
正文。

# 2. 文献综述

## 2.1 理论基础

### 2.1.1 绿色全要素生产率的内涵与测度
正文。
""",
        encoding="utf-8",
    )
    output = tmp_path / "heading3.docx"
    doc = docx.Document()
    doc.add_heading("摘要", level=1)
    doc.add_paragraph("摘要正文。")
    doc.add_heading("1. 引言", level=1)
    doc.add_paragraph("正文。")
    doc.add_heading("2. 文献综述", level=1)
    doc.add_heading("2.1 理论基础", level=2)
    doc.add_heading("2.1.1 绿色全要素生产率的内涵与测度", level=3)
    doc.add_paragraph("正文。")
    doc.save(output)

    manifest = build_document_structure_manifest(md.read_text(encoding="utf-8"), "zh")
    (tmp_path / "document_structure_manifest.json").write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
    monkeypatch.setattr(post, "_update_fields_with_libreoffice", lambda _path, stats: False)

    stats = post.insert_academic_structure(output, options={"language": "zh", "title": "Heading 3 TOC"})
    report = final_artifact_validation(
        md,
        output,
        tmp_path / "document_structure_manifest.json",
        output_dir=tmp_path,
        telemetry={"toc": {"refreshed": False}},
    )

    assert stats["inspection"]["includes_heading_3"] is True
    assert report["toc"]["field_inserted"] is True
    assert report["toc"]["depth"] == 3
    assert report["toc"]["includes_heading_3"] is True
    assert report["headings"]["literature_review_heading_3_count"] == 1


def test_docx_postprocess_does_not_create_new_conclusion_heading3(monkeypatch, tmp_path):
    docx = pytest.importorskip("docx")
    import utils.docx_post_processor as post

    md = tmp_path / "conclusion.md"
    md.write_text(
        """---
title: "结论结构测试"
language: "zh"
date: "May 2026"
---

# 6. 结论

## 6.1 研究总结与管理启示

### 6.1.1 核心研究发现的归纳
正文。
""",
        encoding="utf-8",
    )
    output = tmp_path / "conclusion.docx"
    doc = docx.Document()
    doc.add_heading("6. 结论", level=1)
    doc.add_heading("6.1 研究总结与管理启示", level=2)
    para = doc.add_paragraph()
    run = para.add_run("核心研究发现的归纳")
    run.bold = True
    doc.add_paragraph("正文。")
    doc.save(output)

    manifest = build_document_structure_manifest(md.read_text(encoding="utf-8"), "zh")
    (tmp_path / "document_structure_manifest.json").write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
    monkeypatch.setattr(post, "_update_fields_with_libreoffice", lambda _path, stats: False)

    post.insert_academic_structure(output, options={"language": "zh", "title": "结论结构测试"})

    processed = docx.Document(output)
    matching = [para for para in processed.paragraphs if para.text.strip() == "核心研究发现的归纳"]
    assert matching
    assert all(not (para.style.name if para.style else "").startswith("Heading") for para in matching)


def test_toc_manual_update_status(monkeypatch, tmp_path):
    docx = pytest.importorskip("docx")
    import utils.docx_post_processor as post

    output = tmp_path / "toc_manual.docx"
    doc = docx.Document()
    doc.add_heading("1. Introduction", level=1)
    doc.add_heading("1.1 Background", level=2)
    doc.add_heading("2. Literature Review", level=1)
    doc.add_heading("2.1 Prior Work", level=2)
    doc.save(output)
    monkeypatch.setattr(post, "_update_fields_with_libreoffice", lambda _path, stats: False)
    monkeypatch.setenv("TOC_STRICT", "true")

    stats = post.insert_academic_structure(output, options={"language": "en", "title": "Manual TOC"})

    assert stats["toc_field_inserted"] is True
    assert stats["toc_refreshed"] is False
    assert stats["manual_update_required"] is True
    assert stats["format_status"] == "incomplete_toc_refresh"


def test_technical_tokens_survive_docx(tmp_path):
    docx = pytest.importorskip("docx")

    md = tmp_path / "tokens.md"
    md.write_text(
        """---
title: "Token Contract"
language: "en"
date: "May 2026"
---

# 1. Introduction

A* D* CCD* D* Lite C++ C# F# R-I X-Y STM32F4 STM32G0 Co²⁺ BO₃ BO₄ SiO₂ B₂O₃
""",
        encoding="utf-8",
    )
    output = tmp_path / "tokens.docx"
    doc = docx.Document()
    _ensure_paragraph_style(doc, "CoverTitle")
    doc.add_paragraph("Token Contract", style="CoverTitle")
    doc.add_heading("1. Introduction", level=1)
    doc.add_paragraph("A* D* CCD* D* Lite C++ C# F# R-I X-Y STM32F4 STM32G0 Co²⁺ BO₃ BO₄ SiO₂ B₂O₃")
    doc.save(output)

    report = final_artifact_validation(
        md,
        output,
        tmp_path / "missing_manifest.json",
        output_dir=tmp_path,
        telemetry={"toc": {"field_inserted": False, "refreshed": False}, "cleanup": {}},
    )

    assert report["technical_tokens"]["damaged_tokens"] == []


def test_format_warnings_matches_final_artifacts(tmp_path):
    docx = pytest.importorskip("docx")

    md = tmp_path / "broken.md"
    md.write_text(
        """---
title: "Broken Contract"
language: "en"
date: "May 2026"
---

# 1. Introduction

Residual {{cite_001}} and {(Author, 2025)}.

<!-- PAGEBREAK --><!-- PAGEBREAK -->

D* appears in source.
""",
        encoding="utf-8",
    )
    output = tmp_path / "broken.docx"
    doc = docx.Document()
    _ensure_paragraph_style(doc, "CoverTitle")
    doc.add_paragraph("Broken Contract", style="CoverTitle")
    doc.add_heading("1. Introduction", level=1)
    doc.add_paragraph("Residual {{cite_001}} and {(Author, 2025)}.")
    doc.add_paragraph("D appears in source.")
    doc.save(output)

    report = final_artifact_validation(
        md,
        output,
        tmp_path / "missing_manifest.json",
        output_dir=tmp_path,
        telemetry={"cleanup": {"duplicate_pagebreaks_fixed": 1, "citation_residuals_fixed": 0}},
    )
    persisted = json.loads((tmp_path / "format_warnings.json").read_text(encoding="utf-8"))

    assert "Final Markdown still contains citation residue." in persisted["warnings_remaining"]
    assert "Final Markdown still contains duplicate pagebreak markers." in persisted["warnings_remaining"]
    assert persisted["cleanup"]["duplicate_pagebreaks_fixed"] == 0
    assert persisted["cleanup"]["citation_residuals_fixed"] == 0
    assert persisted["citation"]["residuals_detected"] is True
    assert persisted["citation"]["residuals_fixed"] == 0
    assert persisted["auto_fixed"] == []
    assert persisted["technical_tokens"]["damaged_tokens"] == ["D*"]
    assert report["technical_tokens"]["damaged_tokens"] == ["D*"]


def test_document_ast_manifest_contract_for_chinese_research_paper():
    sample = """---
title: 通用结构测试
language: zh-cn
---

# 摘要
摘要正文。

# 1. 引言
正文。

## 1.1 背景
正文。

### 1.1.1 细分背景
正文。

# 参考文献
条目。
"""
    normalized, doc = normalize_document_markdown(sample, "中文")
    manifest = build_document_structure_manifest(normalized, "zh-cn")

    assert doc.language == "zh"
    assert "### 1.1.1 细分背景" in normalized
    assert manifest["language"] == "zh"
    assert manifest["toc_depth"] == 3
    assert manifest["front_matter"][0]["word_style"] == "AbstractTitle"
    assert manifest["front_matter"][0]["toc_level"] == 1
    assert manifest["back_matter"][0]["word_style"] == "ReferencesTitle"
    assert any(item["word_style"] == "Heading 3" for item in manifest["body_sections"])


def test_document_ast_caps_deep_headings_without_flattening():
    sample = """# 1. Introduction
Text.

## 1.1 Background
Text.

### 1.1.1 Prior Work
Text.

#### 1.1.1.1 Too Deep
Text.

##### 1.1.1.1.1 Much Too Deep
Text.
"""
    normalized, doc = normalize_document_markdown(sample, "en")

    assert "## 1.1 Background" in normalized
    assert "### 1.1.1 Prior Work" in normalized
    assert "####" not in normalized
    assert "##### " not in normalized
    assert "**Too Deep**" in normalized
    assert any(w["type"] == "heading_too_deep" for w in doc.warnings)


def test_document_ast_flattening_regression_preserves_heading_three():
    sample = """# 2. Literature Review

## 2.1 Theory
Body.

### 2.1.1 Stream A
Body.

### 2.1.2 Stream B
Body.
"""
    normalized, _doc = normalize_document_markdown(sample, "en")

    assert "## 2.1 Theory" in normalized
    assert "### 2.1.1 Stream A" in normalized
    assert "### 2.1.2 Stream B" in normalized
    assert sum(1 for line in normalized.splitlines() if line.startswith("## ")) == 1


def test_literature_review_heading3_preserved():
    from phases.compile import normalize_main_body_headings_for_zh

    sample = """## 2.1 文献综述
### 2.1.1 理论基础与概念框架

**绿色全要素生产率的内涵与测度**
正文。

**人工智能驱动生产率提升的理论机制**
正文。
"""
    normalized = normalize_main_body_headings_for_zh(sample)

    assert "# 2. 文献综述" in normalized
    assert "## 2.1 理论基础与概念框架" in normalized
    assert "### 2.1.1 绿色全要素生产率的内涵与测度" in normalized
    assert "### 2.1.2 人工智能驱动生产率提升的理论机制" in normalized
    assert "**绿色全要素生产率的内涵与测度**" not in normalized


def test_bold_subheading_promoted_to_heading3():
    sample = """# 摘要

**研究问题与方法：** 摘要标签不得转换。

**关键词：** 人工智能；生产率

# 1. 引言

**研究背景与问题提出**
正文。

# 2. 文献综述

## 2.1 理论基础

**绿色全要素生产率的内涵与测度**
正文。

**注：** 表格说明不得转换。
"""
    normalized, doc = normalize_document_markdown(sample, "zh")

    assert "### 2.1.1 绿色全要素生产率的内涵与测度" in normalized
    assert "**研究背景与问题提出**" in normalized
    assert "**关键词：**" in normalized
    assert "**注：**" in normalized
    assert any(section.word_style == "Heading 3" for section in doc.sections)


def test_heading_hierarchy_not_flattened():
    sample = """# 2. 文献综述

## 2.1 文献综述
正文。

### 2.1.1 A
正文。

### 2.1.2 B
正文。
"""
    normalized, doc = normalize_document_markdown(sample, "zh")

    assert "## 2.1 文献综述" in normalized
    assert "### 2.1.1 A" in normalized
    assert "### 2.1.2 B" in normalized
    assert sum(1 for line in normalized.splitlines() if line.startswith("## ")) == 1
    assert any(section.word_style == "Heading 3" for section in doc.sections)


def test_empty_heading_removed_or_downgraded():
    sample = """# 2. 文献综述

## 2.1

## 2.2 下一节
正文。
"""
    normalized, doc = normalize_document_markdown(sample, "zh")

    assert "## 2.1 2.1" not in normalized
    assert "## 2.1\n" not in normalized
    assert "下一节" in normalized
    assert any(w["type"] == "empty_heading" for w in doc.warnings)


def test_docx_post_processor_english_toc_and_body_indent(monkeypatch, tmp_path):
    docx = pytest.importorskip("docx")
    import utils.docx_post_processor as post

    output = tmp_path / "en_toc_indent.docx"
    doc = docx.Document()
    doc.add_heading("Abstract", level=1)
    doc.add_paragraph("Keywords: governance; performance")
    doc.add_heading("1. Introduction", level=1)
    body = doc.add_paragraph("This is an English body paragraph.")
    doc.add_heading("1.1 Background", level=2)
    doc.add_paragraph("More English body text.")
    doc.add_heading("2. Literature Review", level=1)
    doc.add_paragraph("Prior work.")
    doc.save(output)

    monkeypatch.setattr(post, "_update_fields_with_libreoffice", lambda _path, stats: False)

    stats = post.insert_academic_structure(output, options={"language": "en", "title": "English TOC Test"})
    processed = docx.Document(output)
    texts = [p.text.strip() for p in processed.paragraphs if p.text.strip()]

    assert "Table of Contents" in texts
    assert "1. Introduction" in texts
    assert "1.1 Background" in texts
    assert any("TOC field inserted but automatic refresh failed" in warning for warning in stats["warnings"])

    normal_indent = processed.styles["Normal"].paragraph_format.first_line_indent
    assert normal_indent is None or normal_indent.pt == 0
    body_paragraphs = [p for p in processed.paragraphs if p.text.strip() == body.text]
    assert body_paragraphs
    direct_indent = body_paragraphs[0].paragraph_format.first_line_indent
    assert direct_indent is None or direct_indent.pt == 0


def test_docx_post_processor_chinese_toc_and_body_indent(monkeypatch, tmp_path):
    docx = pytest.importorskip("docx")
    import utils.docx_post_processor as post

    output = tmp_path / "zh_toc_indent.docx"
    doc = docx.Document()
    doc.add_heading("摘要", level=1)
    doc.add_paragraph("关键词：治理；绩效")
    doc.add_heading("1. 引言", level=1)
    body = doc.add_paragraph("这是中文正文段落。")
    doc.add_heading("1.1 一、研究背景", level=2)
    doc.add_paragraph("更多中文正文。")
    doc.add_heading("2. 文献综述", level=1)
    doc.add_paragraph("已有研究。")
    doc.save(output)

    monkeypatch.setattr(post, "_update_fields_with_libreoffice", lambda _path, stats: False)

    stats = post.insert_academic_structure(output, options={"language": "zh", "title": "中文目录测试"})
    processed = docx.Document(output)
    texts = [p.text.strip() for p in processed.paragraphs if p.text.strip()]

    assert "目录" in texts
    assert "1. 引言" in texts
    assert "1.1 研究背景" in texts
    assert "1.1 一、研究背景" not in texts
    assert any("TOC field inserted but automatic refresh failed" in warning for warning in stats["warnings"])

    normal_indent = processed.styles["Normal"].paragraph_format.first_line_indent
    assert normal_indent is not None and normal_indent.pt > 0
    body_paragraphs = [p for p in processed.paragraphs if p.text.strip() == body.text]
    assert body_paragraphs
    direct_indent = body_paragraphs[0].paragraph_format.first_line_indent
    assert direct_indent is not None and direct_indent.pt > 0


def test_docx_postprocess_does_not_change_outline(monkeypatch, tmp_path):
    docx = pytest.importorskip("docx")
    import utils.docx_post_processor as post

    output = tmp_path / "outline_guard.docx"
    doc = docx.Document()
    doc.add_heading("1. 引言", level=1)
    doc.add_paragraph("引言正文。")
    doc.add_heading("2. 文献综述", level=1)
    doc.add_heading("2.1 理论基础", level=2)
    doc.add_heading("2.1.1 机制分析", level=3)
    doc.add_heading("3. 研究方法", level=1)
    doc.add_heading("4. 分析结果", level=1)
    doc.add_heading("5. 讨论", level=1)
    doc.add_heading("6. 结论", level=1)
    doc.save(output)

    before = extract_docx_outline(output)
    monkeypatch.setattr(post, "_update_fields_with_libreoffice", lambda _path, stats: False)

    stats = post.insert_academic_structure(output, options={"language": "zh", "title": "结构冻结测试"})
    after = extract_docx_outline(output)

    assert stats["post_processor_success"] is True
    assert outline_signature(before) == outline_signature(after)
    assert [item["number"] for item in after if item["level"] == 1] == ["1", "2", "3", "4", "5", "6"]


def test_docx_post_processor_places_chinese_toc_after_abstract(monkeypatch, tmp_path):
    docx = pytest.importorskip("docx")
    import utils.docx_post_processor as post

    output = tmp_path / "zh_toc_order.docx"
    doc = docx.Document()
    doc.add_paragraph("Table of Contents")
    doc.add_heading("摘要", level=1)
    doc.add_paragraph("摘要正文。")
    doc.add_heading("1. 引言", level=1)
    doc.add_heading("1.1 研究背景", level=2)
    doc.add_paragraph("正文。")
    doc.add_heading("2. 文献综述", level=1)
    doc.add_heading("2.1 理论基础", level=2)
    doc.add_paragraph("正文。")
    doc.save(output)

    monkeypatch.setattr(post, "_update_fields_with_libreoffice", lambda _path, stats: False)

    stats = post.insert_academic_structure(
        output,
        options={
            "language": "zh",
            "title": "中文目录测试",
            "toc_entries": [(1, "1. 引言"), (2, "1.1 研究背景"), (1, "2. 文献综述"), (2, "2.1 理论基础")],
        },
    )
    processed = docx.Document(output)
    texts = [p.text.strip() for p in processed.paragraphs if p.text.strip()]

    assert "Table of Contents" not in texts
    assert texts.count("目录") == 1
    assert texts.index("摘要") < texts.index("目录") < texts.index("1. 引言")
    assert "1.1 研究背景" in texts
    assert any("TOC field inserted but automatic refresh failed" in warning for warning in stats["warnings"])


def test_docx_post_processor_does_not_treat_sentence_as_caption(tmp_path):
    docx = pytest.importorskip("docx")
    from utils.docx_post_processor import insert_academic_structure

    output = tmp_path / "caption_sentence.docx"
    doc = docx.Document()
    doc.add_heading("1. Introduction", level=1)
    doc.add_paragraph("Table 1 shows that the model improves accuracy across settings.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    table.cell(1, 0).text = "x"
    table.cell(1, 1).text = "y"
    doc.save(output)

    stats = insert_academic_structure(output, options={"language": "en", "title": "Caption sentence"})
    processed = docx.Document(output)
    texts = "\n".join(p.text for p in processed.paragraphs)
    assert "Table 1 shows that the model improves accuracy across settings." in texts
    assert stats["captions_generated"] == 0


@pytest.mark.skipif(shutil.which("pandoc") is None, reason="Pandoc is not installed")
def test_docx_export_regression_chinese_drone_sample(tmp_path):
    docx = pytest.importorskip("docx")
    input_md = tmp_path / "drone.md"
    output = tmp_path / "drone.docx"
    input_md.write_text(ZH_DRONE_SAMPLE, encoding="utf-8")

    assert export_docx(input_md, output)

    document = docx.Document(output)
    texts = "\n".join(p.text for p in document.paragraphs)

    forbidden = [
        "Right-click and update field",
        "Table of Contents",
        "Title",
        "Document Type",
        "Disclaimer",
        "Table 1. Table 1",
        "| 技术类型 |",
        "| ---",
        "| Performance |",
        "ewpage",
        "\\newpage",
        "{(",
        ")}",
    ]
    for marker in forbidden:
        assert marker not in texts

    required = ["目录", "摘要", "关键词", "1. 引言", "2. 正文", "3. 结论", "参考文献"]
    for marker in required:
        assert marker in texts

    for para in document.paragraphs:
        style = para.style.name if para.style else ""
        if style.startswith("Heading") and para.text.strip():
            assert not para.text.strip().startswith(("·", "-", "*"))

    assert len(document.tables) > 0
    captions = [p.text.strip() for p in document.paragraphs if p.text.strip().startswith(("表 ", "Table "))]
    assert captions
    assert all(not caption.startswith("Table ") for caption in captions)
    assert len(captions) == len(set(captions))
    numbers = [caption.split()[1] for caption in captions if caption.startswith("表 ") and len(caption.split()) > 1]
    assert len(numbers) == len(set(numbers))


def test_docx_post_processor_rejects_lost_markdown_tables(tmp_path):
    docx = pytest.importorskip("docx")
    from utils.docx_post_processor import insert_academic_structure

    output = tmp_path / "lost_table.docx"
    doc = docx.Document()
    doc.add_heading("1. 引言", level=1)
    doc.add_paragraph("表 1 | A | B |")
    doc.save(output)

    stats = insert_academic_structure(
        output,
        options={"language": "zh", "title": "表格失败测试", "markdown_tables_detected": 1},
    )

    assert stats["post_processor_success"] is True
    assert stats["validation_errors"]
    assert any("Markdown tables were detected before export" in warning for warning in stats["validation_errors"])
