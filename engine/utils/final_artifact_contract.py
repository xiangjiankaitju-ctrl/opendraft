#!/usr/bin/env python3
"""Final Markdown/DOCX artifact contracts and validation helpers."""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
import copy
import json
import re
from typing import Any, Optional


PAGEBREAK_MARKER = "<!-- PAGEBREAK -->"

TECHNICAL_TOKENS: tuple[str, ...] = (
    "D* Lite",
    "CCD*",
    "A*",
    "D*",
    "C++",
    "C#",
    "F#",
    "R-I",
    "X-Y",
    "STM32F4",
    "STM32G0",
    "Co²⁺",
    "BO₃",
    "BO₄",
    "SiO₂",
    "B₂O₃",
)

_TECHNICAL_TOKEN_PATTERNS: dict[str, re.Pattern[str]] = {
    "D* Lite": re.compile(r"(?<![A-Za-z0-9\\])D\\?\*\s+Lite\b"),
    "CCD*": re.compile(r"(?<![A-Za-z0-9\\])CCD\\?\*(?![A-Za-z0-9])"),
    "A*": re.compile(r"(?<![A-Za-z0-9\\])A\\?\*(?![A-Za-z0-9])"),
    "D*": re.compile(r"(?<![A-Za-z0-9\\])D\\?\*(?!\s+Lite\b)(?![A-Za-z0-9])"),
    "C++": re.compile(r"\bC\+\+"),
    "C#": re.compile(r"\bC#"),
    "F#": re.compile(r"\bF#"),
    "R-I": re.compile(r"\bR-I\b"),
    "X-Y": re.compile(r"\bX-Y\b"),
    "STM32F4": re.compile(r"\bSTM32F4\b"),
    "STM32G0": re.compile(r"\bSTM32G0\b"),
    "Co²⁺": re.compile(r"Co²⁺"),
    "BO₃": re.compile(r"BO₃"),
    "BO₄": re.compile(r"BO₄"),
    "SiO₂": re.compile(r"SiO₂"),
    "B₂O₃": re.compile(r"B₂O₃"),
}


@dataclass
class Metadata:
    """Structured final Markdown metadata."""

    title: str = ""
    author: str = "OpenDraft AI"
    date: str = ""
    language: str = "en"
    extras: dict[str, Any] = field(default_factory=dict)


def normalize_language(value: Any) -> str:
    raw = str(value or "").strip().lower()
    if raw in {"zh", "zh-cn", "zh_cn", "chinese", "cn", "中文"}:
        return "zh"
    if raw in {"en", "english", "英文"}:
        return "en"
    return "en"


def yaml_quote(value: Any) -> str:
    return str(value or "").replace("\\", "\\\\").replace('"', '\\"')


def write_front_matter(metadata: Metadata) -> str:
    """Serialize final YAML front matter. This is the single YAML writer."""
    language = normalize_language(metadata.language)
    date = str(metadata.date or datetime.now().strftime("%B %Y")).strip()
    title = str(metadata.title or "").strip()
    author = str(metadata.author or "OpenDraft AI").strip() or "OpenDraft AI"
    lines = [
        f'title: "{yaml_quote(title)}"',
        f'author: "{yaml_quote(author)}"',
        f'date: "{yaml_quote(date)}"',
        f'language: "{language}"',
    ]
    reserved = {"title", "author", "date", "language", "lang"}
    for key, value in metadata.extras.items():
        key_text = str(key or "").strip()
        if not key_text or key_text.lower() in reserved:
            continue
        lines.append(f'{key_text}: "{yaml_quote(value)}"')
    return "---\n" + "\n".join(lines) + "\n---"


def split_front_matter(content: str) -> tuple[list[str], str] | None:
    if not content.lstrip().startswith("---"):
        return None
    leading = len(content) - len(content.lstrip())
    if leading:
        content = content.lstrip()
    parts = content.split("---", 2)
    if len(parts) < 3:
        return None
    return parts[1].splitlines(), parts[2]


def extract_front_matter_metadata(content: str) -> dict[str, str]:
    split = split_front_matter(content)
    if split is None:
        return {}
    lines, _body = split
    data: dict[str, str] = {}
    for line in lines:
        if ":" not in line or line.lstrip().startswith("#"):
            continue
        key, value = line.split(":", 1)
        key = key.strip()
        if key:
            data[key] = value.strip().strip("'\"")
    return data


def validate_front_matter_schema(
    content: str,
    language: Any = None,
    *,
    title_candidates: Optional[list[tuple[str, Any]]] = None,
    filename_fallback: str = "research_paper",
) -> tuple[str, dict[str, Any]]:
    """Validate and repair final Markdown front matter from structured metadata."""
    requested_language = normalize_language(language) if str(language or "").strip() else ""
    split = split_front_matter(content)
    repaired_missing_title = False
    repaired_keyless_title = False
    repaired_language = False
    repaired_date = False

    if split is None:
        lines: list[str] = []
        body = content
        repaired_missing_title = True
    else:
        lines, body = split

    existing: dict[str, str] = {}
    extras: dict[str, Any] = {}
    keyless_title = ""
    for line in lines:
        keyless = re.match(r'^\s*:\s*["\']?(.+?)["\']?\s*$', line)
        if keyless:
            keyless_title = keyless.group(1).strip()
            repaired_keyless_title = True
            continue
        if ":" not in line or line.lstrip().startswith("#"):
            continue
        key, value = line.split(":", 1)
        key = key.strip()
        if not key:
            continue
        clean_value = value.strip().strip("'\"")
        existing[key] = clean_value
        if key.lower() not in {"title", "author", "date", "language", "lang"}:
            extras[key] = clean_value

    if str(existing.get("title") or "").strip():
        title = str(existing.get("title") or "").strip()
        title_source = "metadata.title"
    elif keyless_title:
        title = keyless_title
        title_source = "keyless_yaml"
    else:
        title = ""
        title_source = ""
        candidates = list(title_candidates or [])
        candidates.append(("filename", filename_fallback))
        for source, value in candidates:
            candidate = str(value or "").strip()
            if candidate:
                title = candidate
                title_source = source
                break
    if not title:
        title = "research_paper"
        title_source = "default"
    if not existing.get("title"):
        repaired_missing_title = True

    old_language = str(existing.get("language") or existing.get("lang") or "").strip().lower()
    final_language = requested_language or normalize_language(old_language)
    if old_language not in {"zh", "en"} or old_language != final_language:
        repaired_language = True

    date_value = str(existing.get("date") or "").strip()
    if not date_value:
        repaired_date = True
        date_value = datetime.now().strftime("%B %Y")

    metadata = Metadata(
        title=title,
        author=existing.get("author") or "OpenDraft AI",
        date=date_value,
        language=final_language,
        extras=extras,
    )
    fixed = write_front_matter(metadata) + body
    final_metadata = extract_front_matter_metadata(fixed)
    front_yaml = split_front_matter(fixed)[0] if split_front_matter(fixed) else []
    has_keyless_yaml = any(re.match(r'^\s*:\s*["\']?.+?["\']?\s*$', line) for line in front_yaml)
    front_matter_valid = (
        bool(final_metadata.get("title"))
        and final_metadata.get("language") in {"zh", "en"}
        and bool(final_metadata.get("date"))
        and not has_keyless_yaml
    )
    report = {
        "title_present": bool(final_metadata.get("title")),
        "title_source": title_source,
        "front_matter_valid": front_matter_valid,
        "repaired_missing_title": repaired_missing_title,
        "repaired_keyless_title": repaired_keyless_title,
        "repaired_language": repaired_language,
        "repaired_date": repaired_date,
    }
    return fixed, report


def normalize_pagebreaks(text: str) -> str:
    """Normalize all supported pagebreak variants to the final marker."""
    if not text:
        return text
    normalized = text
    htmlish = r"<!\s*(?:--|—|–|-)+\s*PAGEBREAK\s*(?:--|—|–|-)+\s*>"
    normalized = re.sub(htmlish, PAGEBREAK_MARKER, normalized, flags=re.IGNORECASE)
    normalized = re.sub(r"<!--\s*PAGEBREAK\s*-->", PAGEBREAK_MARKER, normalized, flags=re.IGNORECASE)
    normalized = re.sub(
        r"(?im)^[ \t]*(?:\\+newpage|/newpage|ewpage|newpage)[ \t]*" + re.escape(PAGEBREAK_MARKER) + r"[ \t]*$",
        PAGEBREAK_MARKER,
        normalized,
    )
    normalized = re.sub(
        re.escape(PAGEBREAK_MARKER) + r"[ \t]*(?:\\+newpage|/newpage|ewpage|newpage)[ \t]*",
        PAGEBREAK_MARKER,
        normalized,
        flags=re.IGNORECASE,
    )
    normalized = re.sub(
        r"(?im)^[ \t]*(?:\\+newpage|/newpage|ewpage|newpage)[ \t]*$",
        PAGEBREAK_MARKER,
        normalized,
    )
    normalized = re.sub(
        r"(?is)(?:\s*" + re.escape(PAGEBREAK_MARKER) + r"\s*){2,}",
        f"\n\n{PAGEBREAK_MARKER}\n\n",
        normalized,
    )
    return normalized


def has_duplicate_pagebreaks(text: str) -> bool:
    return bool(re.search(r"(?is)(?:\s*<!--\s*PAGEBREAK\s*-->\s*){2,}", text or ""))


def find_malformed_pagebreaks(text: str) -> list[str]:
    issues: list[str] = []
    for line in (text or "").splitlines():
        if "PAGEBREAK" in line and line.strip() != PAGEBREAK_MARKER:
            issues.append(line.strip())
        elif re.match(r"(?i)^\s*(?:\\+newpage|/newpage|ewpage|newpage)\s*$", line):
            issues.append(line.strip())
    return issues


def clean_citation_residuals(text: str, language: Any = "en") -> str:
    """Remove raw cite IDs and normalize brace-wrapped author-year citations."""
    lang = normalize_language(language)
    out = text
    out = re.sub(r"\{\{\s*cite_\d{3,}\s*\}\}", "", out)
    out = re.sub(r"\{\s*cite_\d{3,}\s*\}", "", out)
    out = re.sub(r"\bcite_\d{3,}\b", "", out)
    author_year = r"[^{}\n]*?(?:\d{4}|n\.d\.)[^{}\n]*?"
    if lang == "zh":
        out = re.sub(r"\{\s*\((" + author_year + r")\)\s*\}", r"（\1）", out)
        out = re.sub(r"\{\{\s*\((" + author_year + r")\)\s*\}\}", r"（\1）", out)
        out = re.sub(r"\{\{\s*(" + author_year + r")\s*\}\}", r"（\1）", out)
        out = re.sub(r"\{\s*(" + author_year + r")\s*\}", r"（\1）", out)
        out = re.sub(r"\(\s*([^()\n]*?,\s*(?:\d{4}|n\.d\.))\s*\)", r"（\1）", out)
        out = re.sub(r"([\u4e00-\u9fff])\s+（", r"\1（", out)
        out = re.sub(r"）\s+([。；，、])", r"）\1", out)
    else:
        out = re.sub(r"\{\s*\((" + author_year + r")\)\s*\}", r"(\1)", out)
        out = re.sub(r"\{\{\s*\((" + author_year + r")\)\s*\}\}", r"(\1)", out)
        out = re.sub(r"\{\{\s*(" + author_year + r")\s*\}\}", r"(\1)", out)
        out = re.sub(r"\{\s*(" + author_year + r")\s*\}", r"(\1)", out)
        out = re.sub(r"（([^（）\n]*?,\s*(?:\d{4}|n\.d\.))）", r"(\1)", out)
    return out


def find_citation_residuals(text: str) -> list[str]:
    patterns = [
        r"\{\{?\s*cite_\d{3,}\s*\}?\}",
        r"\bcite_\d{3,}\b",
        r"\{\s*\([^{}\n]+?(?:\d{4}|n\.d\.)[^{}\n]*?\)\s*\}",
        r"\{\{\s*[^{}\n]+?(?:\d{4}|n\.d\.)[^{}\n]*?\s*\}\}",
    ]
    residuals: list[str] = []
    for pattern in patterns:
        residuals.extend(match.group(0) for match in re.finditer(pattern, text or ""))
    return residuals


def find_technical_tokens(text: str) -> list[str]:
    normalized = (text or "").replace("\\*", "*")
    found: list[str] = []
    for token in TECHNICAL_TOKENS:
        pattern = _TECHNICAL_TOKEN_PATTERNS[token]
        if pattern.search(normalized):
            found.append(token)
    return found


def detect_damaged_technical_tokens(source_text: str, final_text: str) -> list[str]:
    final_normalized = (final_text or "").replace("\\*", "*")
    damaged: list[str] = []
    for token in find_technical_tokens(source_text):
        if not _TECHNICAL_TOKEN_PATTERNS[token].search(final_normalized):
            damaged.append(token)
    return damaged


def final_artifact_validation(
    final_md_path: Path,
    final_docx_path: Path,
    manifest_path: Optional[Path] = None,
    *,
    output_dir: Optional[Path] = None,
    telemetry: Optional[dict[str, Any]] = None,
) -> dict[str, Any]:
    """Validate final artifacts and write the authoritative warnings report."""
    report: dict[str, Any] = copy.deepcopy(telemetry or {})
    warnings = report.setdefault("warnings", [])
    warnings_remaining = report.setdefault("warnings_remaining", [])

    def add_warning(code: str, message: str, action: str = "Review final artifact.") -> None:
        item = {"code": code, "message": message, "action": action}
        if item not in warnings:
            warnings.append(item)

    def add_remaining(message: str) -> None:
        if message not in warnings_remaining:
            warnings_remaining.append(message)

    final_md_path = Path(final_md_path)
    final_docx_path = Path(final_docx_path)
    manifest_path = Path(manifest_path) if manifest_path else None
    output_dir = Path(output_dir) if output_dir else final_docx_path.parent
    fatal = False

    try:
        md_text = final_md_path.read_text(encoding="utf-8")
    except OSError as exc:
        md_text = ""
        fatal = True
        add_remaining(f"Final Markdown could not be read: {exc}")
    if not md_text.strip():
        fatal = True
        add_remaining("Final Markdown is empty.")

    metadata = extract_front_matter_metadata(md_text)
    language = normalize_language(metadata.get("language") or metadata.get("lang"))
    title = str(metadata.get("title") or final_md_path.stem).strip()
    front_lines = split_front_matter(md_text)[0] if split_front_matter(md_text) else []
    keyless_yaml = any(re.match(r'^\s*:\s*["\']?.+?["\']?\s*$', line) for line in front_lines)
    front_matter_valid = bool(metadata.get("title")) and language in {"zh", "en"} and bool(metadata.get("date")) and not keyless_yaml
    metadata_report = report.setdefault("metadata", {})
    metadata_report.update(
        {
            "title_present": bool(metadata.get("title")),
            "title_source": metadata_report.get("title_source") or ("metadata.title" if metadata.get("title") else "filename"),
            "front_matter_valid": front_matter_valid,
            "repaired_keyless_title": bool(metadata_report.get("repaired_keyless_title", False)),
        }
    )
    if keyless_yaml:
        add_warning("warning_high", "Final Markdown front matter contains keyless YAML.", "Do not mark front_matter_valid=true.")
        add_remaining("Final Markdown front matter contains keyless YAML.")
    if not front_matter_valid:
        add_remaining("Final Markdown front matter schema is invalid.")

    malformed_pagebreaks = find_malformed_pagebreaks(md_text)
    duplicate_pagebreaks = has_duplicate_pagebreaks(md_text)
    if malformed_pagebreaks:
        add_warning("warning_high", "Final Markdown contains malformed pagebreak markers.", "Normalize with normalize_pagebreaks before writing.")
        add_remaining("Final Markdown contains malformed pagebreak markers.")
    if duplicate_pagebreaks:
        add_warning("warning_high", "Final Markdown contains duplicate pagebreak markers.", "Normalize with normalize_pagebreaks before writing.")
        add_remaining("Final Markdown still contains duplicate pagebreak markers.")

    md_citation_residuals = find_citation_residuals(md_text)
    if md_citation_residuals:
        add_warning("warning_high", "Final Markdown still contains citation residue.", "Run final citation cleanup before writing.")
        add_remaining("Final Markdown still contains citation residue.")

    manifest: dict[str, Any] = {}
    if manifest_path and manifest_path.exists():
        try:
            manifest = json.loads(manifest_path.read_text(encoding="utf-8")) or {}
            language = normalize_language(manifest.get("language") or language)
        except (OSError, json.JSONDecodeError):
            manifest = {}

    inspection: dict[str, Any] = {}
    doc_text = ""
    try:
        from docx import Document
        from utils.docx_post_processor import inspect_docx_headings

        if not final_docx_path.exists():
            raise FileNotFoundError(final_docx_path)
        doc = Document(final_docx_path)
        inspection = inspect_docx_headings(final_docx_path, language)
        paragraph_texts = [p.text for p in doc.paragraphs]
        table_texts = [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
        doc_text = "\n".join(paragraph_texts + table_texts)
        cover_paragraphs = [
            p.text.strip()
            for p in doc.paragraphs
            if p.text.strip() and (p.style.name if p.style else "") == "CoverTitle"
        ]
        nonempty_first_three = [p.text.strip() for p in doc.paragraphs if p.text.strip()][:3]
        cover_report = report.setdefault("cover", {})
        cover_report["cover_title"] = cover_paragraphs[0] if cover_paragraphs else cover_report.get("cover_title", "")
        cover_report["cover_title_style_present"] = bool(cover_paragraphs)
        cover_report["title_matches_resolved"] = bool(title and title in cover_paragraphs)
        if cover_paragraphs and title not in cover_paragraphs:
            add_warning("warning_high", "DOCX CoverTitle does not match the resolved paper title.", "Repair cover title source.")
            add_remaining("DOCX CoverTitle does not match the resolved paper title.")
        elif not cover_paragraphs and title and any(title in text for text in nonempty_first_three):
            add_warning("warning_low", "DOCX title appears in the first paragraphs but is not styled CoverTitle.", "Apply CoverTitle style to the cover title.")
        elif title and not cover_paragraphs:
            add_warning("warning_high", "DOCX first page has no CoverTitle paragraph for the resolved title.", "Insert CoverTitle before export completion.")
            add_remaining("DOCX first page has no CoverTitle paragraph for the resolved title.")
        if doc_text and "PAGEBREAK" in doc_text:
            add_warning("warning_high", "Final DOCX displays PAGEBREAK text.", "Convert PAGEBREAK markers to Word page breaks before Pandoc export.")
            add_remaining("Final DOCX displays PAGEBREAK text.")
    except Exception as exc:
        fatal = True
        add_remaining(f"Final DOCX could not be opened or inspected: {exc}")

    toc_report = report.setdefault("toc", {})
    field_inserted = bool(inspection.get("toc_field_exists"))
    toc_report["field_inserted"] = field_inserted
    toc_report["depth"] = inspection.get("toc_depth") or toc_report.get("depth") or 3
    toc_report["refreshed"] = bool(toc_report.get("refreshed", False))
    toc_report["manual_update_required"] = bool(field_inserted and not toc_report["refreshed"])
    if field_inserted and not toc_report["refreshed"]:
        toc_report["format_status"] = toc_report.get("format_status") or "needs_manual_toc_update"
    elif field_inserted:
        toc_report["format_status"] = "complete"
    else:
        toc_report["format_status"] = "missing_toc_field"
        add_warning("warning_high", "Final DOCX has no Word TOC field.", "Insert Word TOC field depth 3.")
        add_remaining("Final DOCX has no Word TOC field.")
    if toc_report["depth"] not in {3, None}:
        add_remaining(f"Final DOCX TOC depth is not 3: {toc_report['depth']}")

    heading_counts = inspection.get("heading_count_by_style") or {}
    heading_report = report.setdefault("headings", {})
    heading_report.update(
        {
            "heading_1_count": int(inspection.get("heading_1_count") or 0),
            "heading_2_count": int(inspection.get("heading_2_count") or 0),
            "heading_3_count": int(heading_counts.get("Heading 3") or 0),
        }
    )
    if heading_report["heading_2_count"] > 12 or (heading_report["heading_3_count"] == 0 and heading_report["heading_2_count"] > 8):
        add_warning("warning_high", "Final DOCX heading hierarchy appears flattened.", "Preserve Heading 3 from the Markdown heading tree.")
        add_remaining("Final DOCX heading hierarchy appears flattened.")
    elif heading_report["heading_2_count"] > 8:
        add_warning("warning", "A chapter has many Heading 2 sections; flattening should be reviewed.", "Inspect the Markdown heading tree.")

    docx_citation_residuals = find_citation_residuals(doc_text)
    if docx_citation_residuals:
        add_warning("warning_high", "Final DOCX still contains citation residue.", "Run final citation cleanup before DOCX export.")
        add_remaining("Final DOCX still contains citation residue.")

    damaged_tokens = detect_damaged_technical_tokens(md_text, doc_text)
    technical_report = report.setdefault("technical_tokens", {})
    technical_report["protected"] = bool(find_technical_tokens(md_text))
    technical_report["damaged_tokens"] = damaged_tokens
    if damaged_tokens:
        add_warning("warning_high", f"Final DOCX damaged technical tokens: {damaged_tokens}", "Protect tokens before Pandoc and verify final DOCX text.")
        add_remaining(f"Final DOCX damaged technical tokens: {damaged_tokens}")

    cleanup_report = report.setdefault("cleanup", {})
    if duplicate_pagebreaks or malformed_pagebreaks:
        cleanup_report["duplicate_pagebreaks_fixed"] = 0
    else:
        cleanup_report["duplicate_pagebreaks_fixed"] = int(cleanup_report.get("duplicate_pagebreaks_fixed") or 0)
    if md_citation_residuals or docx_citation_residuals:
        cleanup_report["citation_residuals_fixed"] = max(int(cleanup_report.get("citation_residuals_fixed") or 0), 1)
    else:
        cleanup_report["citation_residuals_fixed"] = int(cleanup_report.get("citation_residuals_fixed") or 0)

    report["docx_inspection"] = inspection
    report["manifest"] = {"available": bool(manifest), "toc_depth": manifest.get("toc_depth")}
    report["fatal"] = fatal
    output_dir.mkdir(parents=True, exist_ok=True)
    (output_dir / "format_warnings.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return report
