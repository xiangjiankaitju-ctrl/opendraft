#!/usr/bin/env python3
"""
DOCX post-processing for production OpenDraft exports.

Pandoc handles the heavy markdown conversion. This module fixes the Word-native
details Pandoc cannot reliably infer from generated academic drafts: language
cover pages, TOC fields, page breaks, captions, tables, styles, and page numbers.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
from typing import Optional, Dict, Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt
from utils.text_utils import normalize_language_code


ZH_EAST_ASIA_FONT = "Noto Serif CJK SC"
ZH_LATIN_FONT = "Times New Roman"
EN_BODY_FONT = "Times New Roman"


def insert_academic_structure(
    docx_path: Path,
    verbose: bool = False,
    options: Optional[Dict[str, Any]] = None,
) -> dict[str, Any]:
    """Apply production DOCX post-processing and return telemetry stats."""
    options = options or {}
    language = normalize_language_code(options.get("language", "en"))
    language = language if language in {"zh", "en"} else "en"
    stats: dict[str, Any] = {
        "tables_processed": 0,
        "docx_tables_detected": 0,
        "captions_generated": 0,
        "warnings": [],
        "validation_errors": [],
    }

    if not docx_path.exists():
        raise FileNotFoundError(f"DOCX file not found: {docx_path}")

    try:
        if verbose:
            print(f"📄 Post-processing DOCX: {docx_path.name}")

        doc = Document(docx_path)
        _configure_document_styles(doc, language)
        _configure_sections(doc, language)
        _clean_visible_residue(doc, language)
        _remove_initial_pandoc_title_block(doc)
        _insert_cover_page(doc, options, language)
        _remove_existing_toc(doc)
        toc_inserted = _ensure_toc(doc, language)
        _normalize_headings(doc, language)
        _normalize_body_paragraph_styles(doc, language)
        _style_captions_and_notes(doc, language, stats)
        _style_tables(doc, language, stats)
        _page_break_before_references(doc, language)
        _page_break_after_abstract(doc, language)
        if toc_inserted:
            _add_page_numbers(doc)
        _validate_no_math_loss(doc)
        _validate_phase_one_docx(
            doc,
            language,
            int(options.get("markdown_tables_detected") or 0),
            stats,
            options.get("markdown_table_column_counts") or [],
        )

        if _keep_docx_debug_artifacts():
            _emit_docx_debug_stats(doc, stats, "postprocessed_before_lo")
            doc.save(docx_path.parent / "postprocessed_before_lo.docx")
        doc.save(docx_path)
        refreshed = _update_fields_with_libreoffice(docx_path, stats)
        if toc_inserted:
            refreshed_doc = Document(docx_path)
            if not refreshed or not _toc_has_minimum_entries(refreshed_doc):
                _remove_existing_toc(refreshed_doc)
                stats["warnings"].append("TOC refresh did not produce at least 2 entries with page numbers; removed TOC block.")
                refreshed_doc.save(docx_path)
        refreshed_doc = Document(docx_path)
        _validate_phase_one_docx(
            refreshed_doc,
            language,
            int(options.get("markdown_tables_detected") or 0),
            stats,
            options.get("markdown_table_column_counts") or [],
        )
        if _keep_docx_debug_artifacts():
            _emit_docx_debug_stats(refreshed_doc, stats, "final")
            shutil.copy2(docx_path, docx_path.parent / "final.docx")

        if verbose:
            print(
                "   ✅ Post-processing complete "
                f"({stats['tables_processed']} tables, {stats['captions_generated']} captions)"
            )
        return stats
    except Exception as exc:
        if verbose:
            print(f"   ❌ Post-processing failed: {exc}")
            import traceback
            traceback.print_exc()
        return {}


def _configure_document_styles(doc: Document, language: str) -> None:
    _ensure_paragraph_style(doc, "Normal", language, size=12, bold=False)
    _ensure_paragraph_style(doc, "Body Text", language, size=12, bold=False)
    _ensure_paragraph_style(doc, "Caption", language, size=10, bold=False)
    _ensure_paragraph_style(doc, "Table Note", language, size=9, bold=False)
    _ensure_paragraph_style(doc, "References", language, size=11, bold=False)
    _ensure_paragraph_style(doc, "Abstract Label", language, size=12, bold=True)
    _ensure_paragraph_style(doc, "Keywords", language, size=12, bold=False)
    for style_name in ("Heading 1", "Heading 2", "Heading 3"):
        if style_name in doc.styles:
            _set_style_font(doc.styles[style_name], language, size={ "Heading 1": 16, "Heading 2": 14, "Heading 3": 12 }[style_name], bold=True)

    for style_name in ("Normal", "Body Text"):
        if style_name in doc.styles:
            pf = doc.styles[style_name].paragraph_format
            pf.line_spacing = 1.5
            pf.space_after = Pt(6)
            pf.first_line_indent = Pt(24 if language == "zh" else 18)


def _set_style_font(style, language: str, size: Optional[int] = None, bold: Optional[bool] = None, italic: Optional[bool] = None) -> None:
    font = style.font
    font.name = ZH_EAST_ASIA_FONT if language == "zh" else EN_BODY_FONT
    if size:
        font.size = Pt(size)
    elif font.size is None:
        font.size = Pt(10.5 if language == "zh" else 11)
    if bold is not None:
        font.bold = bold
    if italic is not None:
        font.italic = italic
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), ZH_LATIN_FONT if language == "zh" else EN_BODY_FONT)
    rfonts.set(qn("w:hAnsi"), ZH_LATIN_FONT if language == "zh" else EN_BODY_FONT)
    rfonts.set(qn("w:eastAsia"), ZH_EAST_ASIA_FONT if language == "zh" else EN_BODY_FONT)


def _set_run_font(run, language: str) -> None:
    run.font.name = ZH_EAST_ASIA_FONT if language == "zh" else EN_BODY_FONT
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), ZH_LATIN_FONT if language == "zh" else EN_BODY_FONT)
    rfonts.set(qn("w:hAnsi"), ZH_LATIN_FONT if language == "zh" else EN_BODY_FONT)
    rfonts.set(qn("w:eastAsia"), ZH_EAST_ASIA_FONT if language == "zh" else EN_BODY_FONT)


def _ensure_paragraph_style(
    doc: Document,
    name: str,
    language: str,
    size: int,
    bold: bool,
    only_if_missing: bool = False,
) -> None:
    if name in doc.styles:
        if only_if_missing:
            return
        style = doc.styles[name]
    else:
        style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    _set_style_font(style, language, size=size, bold=bold)


def _configure_sections(doc: Document, language: str) -> None:
    for section in doc.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(2.8 if language == "zh" else 2.54)
        section.right_margin = Cm(2.8 if language == "zh" else 2.54)


def _clean_visible_residue(doc: Document, language: str) -> None:
    for para in doc.paragraphs:
        text = para.text
        if not text:
            continue
        cleaned = text
        cleaned = re.sub(r"\*\*(.*?)\*\*", r"\1", cleaned)
        cleaned = re.sub(r"\{\s*(\([^{}\n]+?\))\s*\}", r"\1", cleaned)
        cleaned = re.sub(r"(?i)https://doi\.org/", "https://doi.org/", cleaned)
        cleaned = re.sub(r"(?i)http://doi\.org/", "https://doi.org/", cleaned)
        if language == "zh":
            cleaned = cleaned.replace("Table of Contents", "")
            cleaned = cleaned.replace("Research Problem and Approach", "研究问题与研究方法")
            cleaned = cleaned.replace("Methodology and Findings", "研究方法与主要发现")
            cleaned = cleaned.replace("Key Contributions", "主要贡献")
            cleaned = cleaned.replace("Implications", "理论与现实意义")
            cleaned = re.sub(r"\bKeywords\s*[:：]", "关键词：", cleaned)
        cleaned = re.sub(r"Right-click and update field to refresh the table of contents\.", "", cleaned, flags=re.IGNORECASE)
        cleaned = re.sub(r"(?im)^\s*(?:Title|题目\s*[:：]?|Generated by\s*[:：]?.*|生成工具\s*[:：]?.*)\s*$", "", cleaned)
        cleaned = re.sub(r"(?im)^\s*(?:Disclaimer|This content is for research organization and writing reference only\.|使用声明\s*[:：]?.*)\s*$", "", cleaned)
        cleaned = re.sub(r"OpenDraft AI\s*-?\s*https://github\.com/federicodeponte/opendraft", "", cleaned, flags=re.IGNORECASE)
        cleaned = re.sub(r"\{\s*cite_\d{3,}\s*\}|\bcite_\d{3,}\b", "", cleaned)
        cleaned = re.sub(r"\\+newpage|/newpage|\bewpage\b|\bnewpage\b", "", cleaned, flags=re.IGNORECASE)
        if (para.style.name if para.style else "").startswith("Heading"):
            cleaned = re.sub(r"^[·•\-*]\s+(?=\d+(?:\.\d+)*\.?\s+)", "", cleaned.strip())
        if cleaned != text:
            _replace_paragraph_text(para, cleaned)


def _remove_initial_pandoc_title_block(doc: Document) -> None:
    for para in list(doc.paragraphs[:12]):
        style = para.style.name if para.style else ""
        if style in {"Title", "Subtitle", "Author", "Date"}:
            _delete_paragraph(para)


def _insert_cover_page(doc: Document, options: dict[str, Any], language: str) -> None:
    first = doc.paragraphs[0] if doc.paragraphs else doc.add_paragraph()
    title = options.get("title") or _first_nonempty_heading(doc) or ("文稿题目" if language == "zh" else "Title")
    date = options.get("date") or datetime.now().strftime("%Y-%m-%d")
    project_type = options.get("project_type")

    if language == "zh":
        lines = [(title, 20, True)]
        if project_type:
            lines.append((str(project_type), 12, False))
        if date:
            lines.append((_format_chinese_cover_date(str(date)), 11, False))
    else:
        lines = [(title, 18, True)]
        if project_type:
            lines.append((str(project_type), 12, False))
        if date:
            lines.append((str(date), 11, False))

    inserted = []
    for text, size, bold in lines:
        para = first.insert_paragraph_before(text)
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.first_line_indent = None
        para.paragraph_format.space_after = Pt(10 if bold else 6)
        for run in para.runs:
            _set_run_font(run, language)
            run.font.size = Pt(size)
            run.font.bold = bold
        inserted.append(para)
    spacer = first.insert_paragraph_before("")
    spacer.paragraph_format.space_after = Pt(24)
    inserted.append(spacer)
    inserted[-1].add_run().add_break(WD_BREAK.PAGE)


def _format_chinese_cover_date(date_text: str) -> str:
    match = re.match(r"^(\d{4})[-/年](\d{1,2})(?:[-/月]\d{1,2})?", date_text.strip())
    if match:
        year, month = match.groups()
        return f"{year}年{int(month)}月"
    return date_text


def _ensure_toc(doc: Document, language: str) -> bool:
    if not _should_insert_toc(doc, language):
        return False
    toc_title = "目录" if language == "zh" else "Table of Contents"
    existing = [para for para in doc.paragraphs if para.text.strip() in {"Table of Contents", "目录"} or _paragraph_has_toc_field(para)]
    keep = existing[0] if existing else None
    for para in existing[1:]:
        _delete_paragraph(para)
    if keep:
        _replace_paragraph_text(keep, toc_title)
        keep.style = doc.styles["Heading 1"]
        keep.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _remove_numbering_from_paragraph(keep)
        field_para = _insert_paragraph_after(keep, "")
        _append_toc_field(field_para)
        _ensure_page_break_after(field_para)
        return True

    insert_before = _find_first_body_paragraph(doc, language)
    toc_heading = insert_before.insert_paragraph_before(toc_title) if insert_before else doc.add_paragraph(toc_title)
    toc_heading.style = doc.styles["Heading 1"]
    toc_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _remove_numbering_from_paragraph(toc_heading)
    field_para = _insert_paragraph_after(toc_heading, "")
    _append_toc_field(field_para)
    _ensure_page_break_after(field_para)
    return True


def _should_insert_toc(doc: Document, language: str) -> bool:
    has_level_1 = False
    has_level_2 = False
    for para in doc.paragraphs:
        text = para.text.strip()
        style = para.style.name if para.style else ""
        if not text or not style.startswith("Heading"):
            continue
        if text in {"目录", "Table of Contents"}:
            continue
        if style == "Heading 1":
            has_level_1 = True
        if style == "Heading 2":
            has_level_2 = True
    return has_level_1 and has_level_2


def _remove_existing_toc(doc: Document) -> None:
    """Remove visible/field TOC leftovers until real updated TOC support is stable."""
    removing = False
    for para in list(doc.paragraphs):
        text = para.text.strip()
        style = para.style.name if para.style else ""
        if text in {"Table of Contents", "目录"} or _paragraph_has_toc_field(para):
            removing = True
            _delete_paragraph(para)
            continue
        if removing:
            if style == "Heading 1" and text and text not in {"Table of Contents", "目录"}:
                removing = False
                continue
            if not text or style.startswith("TOC") or re.search(r"\t\d+\s*$", text) or re.search(r"\s\d+\s*$", text):
                _delete_paragraph(para)


def _find_toc_title_paragraph(doc: Document):
    for para in doc.paragraphs:
        if para.text.strip() in {"Table of Contents", "目录"} or _paragraph_has_toc_field(para):
            return para if para.text.strip() else _previous_paragraph(doc, para)
    return None


def _paragraph_has_toc_field(para) -> bool:
    xml = para._p.xml
    return "TOC" in xml and ("instrText" in xml or "fldSimple" in xml)


def _toc_has_minimum_entries(doc: Document) -> bool:
    entries = 0
    in_toc = False
    for para in doc.paragraphs:
        text = para.text.strip()
        style = para.style.name if para.style else ""
        if text in {"目录", "Table of Contents"} or _paragraph_has_toc_field(para):
            in_toc = True
            continue
        if not in_toc:
            continue
        if style == "Heading 1" and text:
            break
        if style.startswith("TOC") and _looks_like_toc_entry_with_page(text):
            entries += 1
            continue
        if _looks_like_toc_entry_with_page(text):
            entries += 1
    return entries >= 2


def _looks_like_toc_entry_with_page(text: str) -> bool:
    if not text:
        return False
    if text in {"目录", "Table of Contents"}:
        return False
    return bool(re.search(r"(?:\t| {2,}|\.{2,})\d+\s*$", text) or re.search(r"\b\d+(?:\.\d+)+\s+.+\s+\d+\s*$", text))


def _append_toc_field(para) -> None:
    run = para.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r'TOC \o "1-3" \h \z \u'
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = ""
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, text, fld_end])


def _normalize_headings(doc: Document, language: str) -> None:
    for para in doc.paragraphs:
        text = para.text.strip()
        style = para.style.name if para.style else ""
        if not text:
            continue
        if _is_reference_heading(text, language):
            if language == "zh":
                text = "参考文献"
                _replace_paragraph_text(para, text)
            para.style = doc.styles["Heading 1"]
            _remove_numbering_from_paragraph(para)
            continue
        if not style.startswith("Heading"):
            continue
        clean = re.sub(r"^[·•\-*]\s+(?=\d+(?:\.\d+)*\.?\s+)", "", text)
        clean = _normalize_formal_heading_text(clean, language)
        unnumbered_candidate = re.sub(r"^\d+(?:\.\d+)*\.?\s+", "", clean).strip()
        if _is_unnumbered_heading(unnumbered_candidate, language):
            clean = unnumbered_candidate
            _remove_numbering_from_paragraph(para)
        else:
            _remove_numbering_from_paragraph(para)
        if clean != text:
            _replace_paragraph_text(para, clean)


def _normalize_formal_heading_text(text: str, language: str) -> str:
    titles = (
        {
            "1": "引言",
            "2": "文献综述",
            "3": "研究方法",
            "4": "分析结果",
            "5": "讨论",
            "6": "结论",
        }
        if language == "zh"
        else {
            "1": "Introduction",
            "2": "Literature Review",
            "3": "Methodology",
            "4": "Analysis and Results",
            "5": "Discussion",
            "6": "Conclusion",
        }
    )
    aliases = {
        "main body": "2",
        "body": "2",
        "正文": "2",
        "literature review": "2",
        "文献综述": "2",
        "methodology": "3",
        "methods": "3",
        "研究方法": "3",
        "analysis": "4",
        "results": "4",
        "analysis and results": "4",
        "分析结果": "4",
        "discussion": "5",
        "讨论": "5",
        "conclusion": "6",
        "conclusions": "6",
        "结论": "6",
    }
    numbered = re.match(r"^(\d)\.?\s+(.+?)\s*$", text)
    if numbered:
        top, body = numbered.groups()
        if top in titles and (body.strip().lower() in aliases or body.strip() in titles.values()):
            return f"{top}. {titles[top]}"
    plain = re.sub(r"^\d+(?:\.\d+)*\.?\s+", "", text).strip()
    mapped = aliases.get(plain.lower())
    if mapped:
        return f"{mapped}. {titles[mapped]}"
    return text


def _normalize_body_paragraph_styles(doc: Document, language: str) -> None:
    before_first_heading = True
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style_name = para.style.name if para.style else ""
        if style_name.startswith("Heading"):
            before_first_heading = False
        if before_first_heading:
            continue
        if style_name.startswith("Heading") or style_name in {"Caption", "Table Note", "References"}:
            continue
        if text in {"目录", "Table of Contents", "Disclaimer"} or _paragraph_has_toc_field(para):
            continue
        if re.match(r"^(?:表\s*\d+|Table\s+\d+|注：|Note:)", text, flags=re.IGNORECASE):
            continue
        para.style = doc.styles["Normal"]
        para.paragraph_format.line_spacing = 1.5
        para.paragraph_format.space_after = Pt(6)
        para.paragraph_format.first_line_indent = Pt(24 if language == "zh" else 18)
        for run in para.runs:
            _set_run_font(run, language)
            run.font.size = Pt(12)


def _is_reference_heading(text: str, language: str) -> bool:
    plain = re.sub(r"^\d+(?:\.\d+)*\.?\s+", "", text.strip()).lower()
    return plain in {"references", "bibliography", "参考文献"}


def _is_unnumbered_heading(text: str, language: str) -> bool:
    plain = re.sub(r"^\d+(?:\.\d+)*\.?\s+", "", text.strip()).lower()
    if plain in {"abstract", "table of contents", "references", "bibliography", "appendix", "appendices"}:
        return True
    if plain.startswith("appendix "):
        return True
    return language == "zh" and plain in {"摘要", "目录", "参考文献", "附录"}


def _apply_heading_numbering(para, level: int) -> None:
    """Attach numPr to headings so Word sees multilevel numbering intent."""
    ppr = para._p.get_or_add_pPr()
    num_pr = ppr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        ppr.append(num_pr)
    ilvl = num_pr.find(qn("w:ilvl"))
    if ilvl is None:
        ilvl = OxmlElement("w:ilvl")
        num_pr.append(ilvl)
    ilvl.set(qn("w:val"), str(level - 1))
    num_id = num_pr.find(qn("w:numId"))
    if num_id is None:
        num_id = OxmlElement("w:numId")
        num_pr.append(num_id)
    num_id.set(qn("w:val"), "1")


def _remove_numbering_from_paragraph(para) -> None:
    ppr = para._p.pPr
    if ppr is None:
        return
    num_pr = ppr.find(qn("w:numPr"))
    if num_pr is not None:
        ppr.remove(num_pr)


def _style_captions_and_notes(doc: Document, language: str, stats: dict[str, Any]) -> None:
    caption_seen = set()
    caption_counter = 0
    table_block_indices = _table_block_indices(doc)
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        if language == "zh" and re.fullmatch(r"Table\s+\d+", text, re.IGNORECASE) and _has_nearby_table(doc, para, table_block_indices, radius=2):
            _replace_paragraph_text(para, "")
            continue
        if (
            len(text) < 120
            and _is_caption_paragraph_text(text, language)
            and _has_nearby_table(doc, para, table_block_indices, radius=2)
        ):
            caption_counter += 1
            body = re.sub(r"^(?:表\s*|Table\s+)\d+(?:[-‑–—]\d+)?[.:：]?\s*", "", text, flags=re.IGNORECASE).strip()
            body = re.sub(r"^(?:表\s*|Table\s+)\d+(?:[-‑–—]\d+)?[.:：]?\s*", "", body, flags=re.IGNORECASE).strip()
            if language == "zh":
                new_text = f"表{caption_counter}" + (f"：{body}" if body else "")
            else:
                new_text = f"Table {caption_counter}" + (f". {body}" if body else "")
            _replace_paragraph_text(para, new_text)
            para.style = doc.styles["Caption"]
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
            para.paragraph_format.keep_with_next = True
            para.paragraph_format.space_before = Pt(6)
            para.paragraph_format.space_after = Pt(3)
            caption_seen.add(caption_counter)
        elif text.startswith("注：") or text.startswith("Note:"):
            para.style = doc.styles["Table Note"]
            para.paragraph_format.space_before = Pt(3)
            para.paragraph_format.space_after = Pt(6)
    stats["captions_generated"] = len(caption_seen)


def _is_caption_paragraph_text(text: str, language: str) -> bool:
    match = re.match(
        r"^(?:表\s*|Table\s+)\d+(?:[-‑–—]\d+)?(?:(?P<punct>[.:：])|\s{2,}|\s+)(?P<body>.*)$",
        text,
        re.IGNORECASE,
    )
    if not match:
        return False
    body = (match.group("body") or "").strip()
    if not body:
        return True
    if match.group("punct"):
        return True
    if language == "zh" and re.match(r"^(?:总结|对比|归纳|显示|说明|表明|展示|列出|给出|呈现)了?", body):
        return False
    if language != "zh" and re.match(r"^(?:shows?|summari[sz]es|compares?|lists?|presents?|indicates?)\b", body, re.IGNORECASE):
        return False
    return True


def _body_block_index(doc: Document, element) -> Optional[int]:
    body_children = list(doc.element.body.iterchildren())
    for idx, child in enumerate(body_children):
        if child is element:
            return idx
    return None


def _table_block_indices(doc: Document) -> set[int]:
    indices = set()
    for table in doc.tables:
        idx = _body_block_index(doc, table._tbl)
        if idx is not None:
            indices.add(idx)
    return indices


def _has_nearby_table(doc: Document, para, table_indices: set[int], radius: int) -> bool:
    para_idx = _body_block_index(doc, para._p)
    if para_idx is None:
        return False
    return any(abs(table_idx - para_idx) <= radius for table_idx in table_indices)


def _style_tables(doc: Document, language: str, stats: dict[str, Any]) -> None:
    for table in doc.tables:
        num_cols = _max_table_cell_count(table)
        if num_cols <= 0:
            raise ValueError("DOCX table has zero columns before styling.")
        _rebuild_table_grid(table, num_cols)
        _set_table_width_pct(table, 5000)
        _set_table_borders(table)
        _disable_table_autofit(table)
        table.autofit = False
        if num_cols > 4:
            _mark_landscape_section(table)

        font_size = Pt(9.5 if language == "zh" else 9) if num_cols > 4 else Pt(10 if language == "zh" else 10)
        for row_idx, row in enumerate(table.rows):
            if row_idx == 0:
                _repeat_table_header(row)
            _prevent_row_split(row)
            for cell in row.cells:
                _set_cell_width(cell, _table_cell_width_dxa(num_cols))
                _set_cell_margins(cell, top=80, bottom=80, left=80, right=80)
                _enable_cell_text_wrap(cell)
                if row_idx == 0:
                    _shade_cell(cell, "EDEDED")
                for para in cell.paragraphs:
                    if para.style is None or para.style.name in {"Normal", "Body Text"}:
                        para.style = doc.styles["Normal"]
                    para.paragraph_format.first_line_indent = None
                    para.paragraph_format.space_before = Pt(2)
                    para.paragraph_format.space_after = Pt(2)
                    for run in para.runs:
                        _set_run_font(run, language)
                        run.font.size = font_size
                        run.font.bold = row_idx == 0
        stats["tables_processed"] += 1
    stats["docx_tables_detected"] = len(doc.tables)


def _max_table_cell_count(table) -> int:
    return max((len(row.cells) for row in table.rows), default=0)


def _table_cell_width_dxa(num_cols: int) -> int:
    usable_width = 9072
    return max(900, usable_width // max(num_cols, 1))


def _rebuild_table_grid(table, num_cols: int) -> None:
    tbl = table._tbl
    for grid in list(tbl.findall(qn("w:tblGrid"))):
        tbl.remove(grid)
    grid = OxmlElement("w:tblGrid")
    width = _table_cell_width_dxa(num_cols)
    for _ in range(num_cols):
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    tbl.insert(0, grid)


def _set_cell_width(cell, width: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width))
    tc_w.set(qn("w:type"), "dxa")


def _set_table_width_pct(table, width: int) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width))
    tbl_w.set(qn("w:type"), "pct")
    jc = tbl_pr.find(qn("w:jc"))
    if jc is None:
        jc = OxmlElement("w:jc")
        tbl_pr.append(jc)
    jc.set(qn("w:val"), "center")


def _disable_table_autofit(table) -> None:
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def _enable_cell_text_wrap(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    for no_wrap in list(tc_pr.findall(qn("w:noWrap"))):
        tc_pr.remove(no_wrap)


def _set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        elem = borders.find(qn(f"w:{edge}"))
        if elem is None:
            elem = OxmlElement(f"w:{edge}")
            borders.append(elem)
        elem.set(qn("w:val"), "single")
        elem.set(qn("w:sz"), "4")
        elem.set(qn("w:space"), "0")
        elem.set(qn("w:color"), "666666")


def _repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        tr_pr.append(OxmlElement("w:tblHeader"))


def _prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def _set_cell_margins(cell, top: int, bottom: int, left: int, right: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for edge, value in {"top": top, "bottom": bottom, "left": left, "right": right}.items():
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _mark_landscape_section(table) -> None:
    # Word section wrapping around individual tables is fragile through
    # python-docx. Marking the table with a wide-table property lets tests and
    # downstream repair code identify it, while width/font fixes keep it usable.
    tbl_pr = table._tbl.tblPr
    desc = tbl_pr.find(qn("w:tblDescription"))
    if desc is None:
        desc = OxmlElement("w:tblDescription")
        tbl_pr.append(desc)
    desc.set(qn("w:val"), "wide-table-landscape-candidate")


def _page_break_after_abstract(doc: Document, language: str) -> None:
    in_abstract = False
    last = None
    for para in doc.paragraphs:
        text = para.text.strip()
        style = para.style.name if para.style else ""
        if text in {"Abstract", "摘要"} and style.startswith("Heading"):
            in_abstract = True
            continue
        if in_abstract and style == "Heading 1" and text and text not in {"目录", "Table of Contents"}:
            if last is not None:
                _ensure_page_break_after(last)
            return
        if in_abstract:
            last = para


def _page_break_before_references(doc: Document, language: str) -> None:
    for para in doc.paragraphs:
        if _is_reference_heading(para.text.strip(), language):
            _ensure_page_break_before(para)
            for following in _iter_following_paragraphs(doc, para):
                if not following.text.strip():
                    continue
                if (following.style.name if following.style else "") == "Heading 1":
                    break
                following.style = doc.styles["References"]
            return


def _add_page_numbers(doc: Document) -> None:
    for idx, section in enumerate(doc.sections):
        section.different_first_page_header_footer = True
        first_footer = section.first_page_footer
        first_para = first_footer.paragraphs[0] if first_footer.paragraphs else first_footer.add_paragraph()
        first_para.text = ""
        footer = section.footer
        para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.text = ""
        run = para.add_run()
        fld_begin = OxmlElement("w:fldChar")
        fld_begin.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = "PAGE"
        fld_end = OxmlElement("w:fldChar")
        fld_end.set(qn("w:fldCharType"), "end")
        run._r.extend([fld_begin, instr, fld_end])


def _validate_no_math_loss(doc: Document) -> None:
    text = "\n".join(para.text for para in doc.paragraphs)
    if re.search(r"\bthermal gradients\s*\(\)", text, flags=re.IGNORECASE):
        raise ValueError("Empty thermal gradient placeholder detected in DOCX: thermal gradients ()")
    if re.search(r"\bgrowth velocity\s*\(\)", text, flags=re.IGNORECASE):
        raise ValueError("Empty growth velocity placeholder detected in DOCX: growth velocity ()")
    if re.search(r"\bcooling rates exceeding\s+K/s\b", text, flags=re.IGNORECASE):
        raise ValueError("Missing numeric value before K/s in DOCX")
    if re.search(r"\bresulting\s+ratio\b", text, flags=re.IGNORECASE):
        raise ValueError("Missing variable before ratio in DOCX")


def _validate_phase_one_docx(
    doc: Document,
    language: str,
    markdown_tables_detected: int,
    stats: dict[str, Any],
    markdown_table_column_counts=None,
) -> None:
    stats["docx_tables_detected"] = len(doc.tables)
    errors: list[str] = []
    texts = [para.text.strip() for para in doc.paragraphs if para.text.strip()]
    full_text = "\n".join(texts)

    if markdown_tables_detected > 0 and len(doc.tables) == 0:
        errors.append(
            f"Markdown tables were detected before export ({markdown_tables_detected}), "
            "but the generated DOCX contains no Word tables."
        )
    if markdown_tables_detected and len(doc.tables) != markdown_tables_detected:
        errors.append(f"Markdown/DOCX table count mismatch: markdown={markdown_tables_detected}, docx={len(doc.tables)}.")

    if any(text in {"目录", "Table of Contents"} for text in texts) and not _toc_has_minimum_entries(doc):
        errors.append("DOCX table of contents is empty or missing page-numbered entries.")
    if _should_insert_toc(doc, language) and "PAGE" not in "\n".join(section.footer._element.xml for section in doc.sections):
        if any(text in {"目录", "Table of Contents"} for text in texts):
            errors.append("DOCX page number field is missing.")
    if "PAGE" in "\n".join(section.first_page_footer._element.xml for section in doc.sections):
        errors.append("Cover page footer contains a page number field.")

    for idx, table in enumerate(doc.tables, start=1):
        max_cells = _max_table_cell_count(table)
        grid_cols = _table_grid_col_count(table)
        if max_cells <= 0:
            errors.append(f"DOCX table {idx} has zero columns.")
        if grid_cols <= 0:
            errors.append(f"DOCX table {idx} has empty tblGrid.")
        elif grid_cols != max_cells:
            errors.append(f"DOCX table {idx} tblGrid/gridCol mismatch: grid={grid_cols}, max_cells={max_cells}.")
        if markdown_table_column_counts and idx <= len(markdown_table_column_counts):
            expected_cols = int(markdown_table_column_counts[idx - 1])
            if expected_cols != max_cells:
                errors.append(f"DOCX table {idx} column count mismatch: markdown={expected_cols}, docx={max_cells}.")

    residue_patterns = [
        r"\bewpage\b",
        r"\\+newpage",
        r"\bcite_\d{3,}\b",
        r"\{\s*\([^{}\n]+?\)\s*\}",
        r"(?m)^\s*(?:Title|题目\s*[:：]?|Generated by\s*[:：]?.*|生成工具\s*[:：]?.*|Disclaimer|使用声明\s*[:：]?.*)\s*$",
        r"\bTable\s+\d+\s*\|",
        r"\b表\s*\d+\s*\|",
        r"\bMain Body\b",
    ]
    for pattern in residue_patterns:
        if re.search(pattern, full_text, flags=re.IGNORECASE):
            errors.append(f"Visible export residue remains in DOCX: {pattern}")

    for para in doc.paragraphs:
        text = para.text.strip()
        style = para.style.name if para.style else ""
        if style.startswith("Heading") and re.match(r"^(?:表|Table)\s*\d+", text, re.IGNORECASE):
            errors.append(f"Table caption is still styled as a heading: {text}")
        if style.startswith("Heading") and len(re.findall(r"\d+", text.split()[0] if text.split() else "")) > 3:
            errors.append(f"Heading depth appears to exceed three levels: {text}")
        if language == "zh" and style in {"Normal", "Body Text"} and text and not _is_cover_or_structural_line(text):
            direct_indent = para.paragraph_format.first_line_indent
            style_indent = para.style.paragraph_format.first_line_indent if para.style else None
            if direct_indent is None and style_indent is None:
                errors.append(f"Chinese body paragraph has no first-line indent: {text[:30]}")
                break

    caption_count = len([
        para.text.strip() for para in doc.paragraphs
        if (para.style.name if para.style else "") == "Caption"
        and re.match(r"^(?:表\s*\d+|Table\s+\d+)\b", para.text.strip(), flags=re.IGNORECASE)
    ])
    if caption_count:
        ref_nums = (
            [int(num) for num in re.findall(r"表\s*(\d+)", full_text)]
            if language == "zh"
            else [int(num) for num in re.findall(r"\bTable\s+(\d+)\b", full_text, flags=re.IGNORECASE)]
        )
        invalid_refs = sorted({num for num in ref_nums if num < 1 or num > caption_count})
        if invalid_refs:
            errors.append(f"Table reference number does not match final caption count: {invalid_refs} > {caption_count}")

    general_forbidden = [
        "Generated by",
        "生成工具",
        "OpenDraft AI - https://github.com/federicodeponte/opendraft",
        "This content is for research organization and writing reference only.",
        "使用声明：本内容仅供资料梳理和写作参考，不构成正式学术成果。",
    ]
    for marker in general_forbidden:
        if marker in full_text:
            errors.append(f"Forbidden cover metadata remains in DOCX: {marker}")

    if language == "zh":
        forbidden = [
            "Title",
            "Document Type",
            "Generated by",
            "生成工具",
            "Disclaimer",
            "OPENDRAFT UNIVERSITY",
            "Faculty of Engineering",
            "Department of Computer Science",
        ]
        for marker in forbidden:
            if re.search(rf"(^|\n)\s*{re.escape(marker)}(?:\s*[:：].*)?(\n|$)", full_text, flags=re.IGNORECASE):
                errors.append(f"English cover/template residue remains in Chinese DOCX: {marker}")

    if errors:
        stats["validation_errors"].extend(errors)
        raise ValueError("; ".join(errors))


def _is_cover_or_structural_line(text: str) -> bool:
    return bool(
        re.match(r"^(?:关键词)", text)
        or re.match(r"^(?:表\s*\d+|Table\s+\d+|注：|Note:)", text, flags=re.IGNORECASE)
    )


def _update_fields_with_libreoffice(docx_path: Path, stats: dict[str, Any]) -> bool:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        stats["warnings"].append("LibreOffice not found; Word fields were inserted but not refreshed server-side.")
        return False
    before = docx_path.with_suffix(".before-lo-refresh.docx")
    shutil.copy2(docx_path, before)
    try:
        with tempfile.TemporaryDirectory() as tmp:
            tmp_dir = Path(tmp)
            working = tmp_dir / docx_path.name
            shutil.copy2(docx_path, working)
            result = subprocess.run(
                [
                    soffice,
                    "--headless",
                    "--convert-to",
                    "docx",
                    "--outdir",
                    str(tmp_dir),
                    str(working),
                ],
                capture_output=True,
                text=True,
                timeout=45,
            )
            if result.returncode != 0:
                stats["warnings"].append(f"LibreOffice field update failed: {result.stderr.strip() or result.returncode}")
                return False
            refreshed = tmp_dir / working.name
            if not refreshed.exists():
                stats["warnings"].append("LibreOffice field update produced no refreshed DOCX.")
                return False
            refreshed_doc = Document(refreshed)
            grid_errors = _table_grid_validation_errors(refreshed_doc)
            if grid_errors:
                stats["warnings"].append("LibreOffice refresh dropped table grid data; keeping pre-refresh DOCX.")
                stats["warnings"].extend(grid_errors)
                return False
            shutil.copy2(refreshed, docx_path)
            return True
    except Exception as exc:
        stats["warnings"].append(f"LibreOffice field update failed: {exc}")
        return False
    finally:
        if before.exists():
            before.unlink()


def render_docx_tables_for_validation(docx_path: Path, stats: dict[str, Any]) -> bool:
    """Render DOCX pages to PNGs when local tools are available."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    pdftoppm = shutil.which("pdftoppm")
    if not soffice or not pdftoppm:
        stats.setdefault("warnings", []).append("DOCX PNG render validation skipped; LibreOffice or pdftoppm is not installed.")
        return False
    try:
        with tempfile.TemporaryDirectory() as tmp:
            tmp_dir = Path(tmp)
            pdf_result = subprocess.run(
                [
                    soffice,
                    "--headless",
                    "--convert-to",
                    "pdf",
                    "--outdir",
                    str(tmp_dir),
                    str(docx_path),
                ],
                capture_output=True,
                text=True,
                timeout=60,
            )
            if pdf_result.returncode != 0:
                stats.setdefault("warnings", []).append(f"DOCX PDF render failed: {pdf_result.stderr.strip() or pdf_result.returncode}")
                return False
            pdf_path = tmp_dir / f"{docx_path.stem}.pdf"
            if not pdf_path.exists():
                stats.setdefault("warnings", []).append("DOCX PDF render produced no PDF.")
                return False
            png_prefix = tmp_dir / "page"
            png_result = subprocess.run(
                [pdftoppm, "-png", str(pdf_path), str(png_prefix)],
                capture_output=True,
                text=True,
                timeout=60,
            )
            if png_result.returncode != 0:
                stats.setdefault("warnings", []).append(f"DOCX PNG render failed: {png_result.stderr.strip() or png_result.returncode}")
                return False
            pngs = sorted(tmp_dir.glob("page-*.png"))
            if not pngs or any(path.stat().st_size < 1000 for path in pngs):
                raise ValueError("DOCX PNG render produced missing or empty page images.")
            stats["rendered_png_pages"] = len(pngs)
            return True
    except Exception as exc:
        stats.setdefault("validation_errors", []).append(f"DOCX PNG render validation failed: {exc}")
        raise


def _table_grid_col_count(table) -> int:
    grid = table._tbl.find(qn("w:tblGrid"))
    if grid is None:
        return 0
    return len(grid.findall(qn("w:gridCol")))


def _table_grid_validation_errors(doc: Document) -> list[str]:
    errors: list[str] = []
    for idx, table in enumerate(doc.tables, start=1):
        max_cells = _max_table_cell_count(table)
        grid_cols = _table_grid_col_count(table)
        if max_cells <= 0 or grid_cols <= 0 or grid_cols != max_cells:
            errors.append(f"Table {idx} grid invalid after refresh: max_cells={max_cells}, grid_cols={grid_cols}")
    return errors


def _keep_docx_debug_artifacts() -> bool:
    import os

    return os.environ.get("OPENDRAFT_KEEP_DOCX_DEBUG", "").strip().lower() in {"1", "true", "yes", "on"}


def _emit_docx_debug_stats(doc: Document, stats: dict[str, Any], stage: str) -> None:
    table_rows = [len(table.rows) for table in doc.tables]
    grid_cols = [_table_grid_col_count(table) for table in doc.tables]
    headings = [
        para.text.strip()
        for para in doc.paragraphs
        if para.text.strip() and (para.style.name if para.style else "").startswith("Heading")
    ]
    stats.setdefault("debug", {})[stage] = {
        "table_count": len(doc.tables),
        "table_row_count": table_rows,
        "grid_column_count": grid_cols,
        "heading_outline": headings,
    }


def _first_nonempty_heading(doc: Document) -> Optional[str]:
    for para in doc.paragraphs:
        text = para.text.strip()
        if text and (para.style.name if para.style else "").startswith("Heading"):
            return re.sub(r"^\d+(?:\.\d+)*\.?\s+", "", text)
    return None


def _find_first_body_paragraph(doc: Document, language: str):
    numbered = re.compile(r"^\s*1\.?\s+(?:引言|Introduction)\s*$", re.IGNORECASE)
    for para in doc.paragraphs:
        text = para.text.strip()
        style = para.style.name if para.style else ""
        if text and style.startswith("Heading") and numbered.match(text):
            return para
    for para in doc.paragraphs:
        text = para.text.strip()
        style = para.style.name if para.style else ""
        if text and style.startswith("Heading") and text not in {"Abstract", "摘要", "Table of Contents", "目录"}:
            return para
    return doc.paragraphs[0] if doc.paragraphs else None


def _find_first_content_paragraph(doc: Document):
    for para in doc.paragraphs:
        text = para.text.strip()
        style = para.style.name if para.style else ""
        if text and style.startswith("Heading") and text not in {"Table of Contents", "目录"}:
            return para
    return doc.paragraphs[0] if doc.paragraphs else None


def _ensure_page_break_after(para) -> None:
    if "w:type=\"page\"" in para._p.xml:
        return
    para.add_run().add_break(WD_BREAK.PAGE)


def _ensure_page_break_before(para) -> None:
    ppr = para._p.get_or_add_pPr()
    if ppr.find(qn("w:pageBreakBefore")) is None:
        ppr.append(OxmlElement("w:pageBreakBefore"))


def _insert_paragraph_after(after_para, text: str):
    new_p = OxmlElement("w:p")
    after_para._p.addnext(new_p)
    from docx.text.paragraph import Paragraph
    return Paragraph(new_p, after_para._parent)


def _delete_paragraph(para) -> None:
    element = para._element
    element.getparent().remove(element)
    para._p = para._element = None


def _replace_paragraph_text(para, text: str) -> None:
    for run in list(para.runs):
        run._element.getparent().remove(run._element)
    if text:
        para.add_run(text)


def _previous_paragraph(doc: Document, para):
    previous = None
    for candidate in doc.paragraphs:
        if candidate._p is para._p:
            return previous
        previous = candidate
    return None


def _iter_following_paragraphs(doc: Document, para):
    found = False
    for candidate in doc.paragraphs:
        if found:
            yield candidate
        elif candidate._p is para._p:
            found = True


def _next_heading_after(doc: Document, para):
    for candidate in _iter_following_paragraphs(doc, para):
        if (candidate.style.name if candidate.style else "").startswith("Heading"):
            return candidate
    return None


def main() -> None:
    import sys

    if len(sys.argv) < 2:
        print("Usage: python docx_post_processor.py <docx-path>")
        raise SystemExit(1)

    result = insert_academic_structure(Path(sys.argv[1]), verbose=True, options={"language": "en"})
    raise SystemExit(0 if result else 1)


if __name__ == "__main__":
    main()
